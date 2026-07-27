#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""発表資料の図表データをExcel化（1図=1シート、数値データ＋ネイティブグラフ）

各シートのデータは 02〜08 の図生成スクリプトと同一ロジックで再計算したもの。
出力: thesis/slides/ゼミ発表_図表データ.xlsx
"""
import os

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(__file__), "..")
PROC = os.path.join(ROOT, "data", "processed")
OUT = os.path.join(ROOT, "slides", "ゼミ発表_図表データ.xlsx")

FONT = "Yu Gothic"
NAVY, BLUE, LBLUE, CORAL, DBLUE = "184F95", "2A78D6", "6DA7EC", "EC835A", "184F95"
RAMP5 = ["CDE2FB", "B7D3F6", "86B6EF", "3987E5", "184F95"]
GRAY = "68758A"
FY_FULL = list(range(2016, 2026))

j = pd.read_csv(os.path.join(PROC, "jepx_kyushu_30min.csv"), parse_dates=["ts"])
daily = pd.read_csv(os.path.join(PROC, "kyushu_daily_metrics.csv"), parse_dates=["day"])
panel = pd.read_csv(os.path.join(PROC, "kyushu_hourly_panel.csv"), parse_dates=["ts"])
bt = pd.read_csv(os.path.join(PROC, "kyushu_battery_backtest_daily.csv"), parse_dates=["day"])
j["hour"] = j["ts"].dt.hour + j["ts"].dt.minute / 60
j["fy"] = j["ts"].dt.year - (j["ts"].dt.month < 4).astype(int)
panel["hour"] = panel["ts"].dt.hour

wb = Workbook()
wb.remove(wb.active)

F_TITLE = Font(name=FONT, size=13, bold=True, color="1F3864")
F_NOTE = Font(name=FONT, size=9, color=GRAY)
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10)
FILL_HDR = PatternFill("solid", fgColor=NAVY)


def new_sheet(name, title, note):
    ws = wb.create_sheet(name)
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = note
    ws["A2"].font = F_NOTE
    ws.sheet_view.showGridLines = True
    return ws


def write_df(ws, df, start_row=4, start_col=1, fmts=None, width=None):
    """DataFrameをヘッダー付きで書き込み。fmts: 列名->番号書式。戻り値: (先頭行, 最終行)"""
    r0 = start_row
    for ci, col in enumerate(df.columns):
        c = ws.cell(row=r0, column=start_col + ci, value=str(col))
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ri, (_, row) in enumerate(df.iterrows()):
        for ci, col in enumerate(df.columns):
            v = row[col]
            if isinstance(v, (np.floating, np.integer)):
                v = v.item()
            c = ws.cell(row=r0 + 1 + ri, column=start_col + ci, value=v)
            c.font = F_BODY
            if fmts and col in fmts:
                c.number_format = fmts[col]
    for ci, col in enumerate(df.columns):
        letter = get_column_letter(start_col + ci)
        ws.column_dimensions[letter].width = (width or {}).get(col, max(10, min(18, len(str(col)) * 2 + 4)))
    return r0, r0 + len(df)


def line_series_colors(chart, colors, smooth=False, marker=True):
    for s, col in zip(chart.series, colors):
        s.smooth = smooth
        s.graphicalProperties.line.solidFill = col
        s.graphicalProperties.line.width = 22000
        if marker:
            s.marker = Marker(symbol="circle", size=5)
            s.marker.graphicalProperties.solidFill = col
            s.marker.graphicalProperties.line.solidFill = col
        else:
            s.marker = Marker(symbol="none")


def style_axes(chart, x_title, y_title):
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.x_axis.delete = False
    chart.y_axis.delete = False


# ============================================================
# f1: 時間帯別平均価格カーブ
# ============================================================
sel = [2016, 2019, 2023, 2025]
f1 = pd.DataFrame({"時刻": sorted(j["hour"].unique())})
for fy in sel:
    g = j[j["fy"] == fy].groupby("hour")["p_kyushu"].mean()
    f1[f"FY{fy}"] = f1["時刻"].map(g)
ws = new_sheet("f1_ダックカーブ", "f1 時間帯別平均価格カーブの変化（ダックカーブの深化）",
               "JEPXスポット九州エリアプライス30分値の年度内平均（円/kWh）。FY2022（燃料危機）は表示対象外。掲載: 本編P4左。再現: analysis/02")
h0, h1 = write_df(ws, f1, fmts={c: "0.00" for c in f1.columns})
ch = LineChart()
ch.title = "時間帯別平均価格カーブ（円/kWh）"
ch.height, ch.width = 11, 20
data = Reference(ws, min_col=2, max_col=1 + len(sel), min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, ["B7D3F6", "6DA7EC", "2A78D6", "184F95"], marker=False)
style_axes(ch, "時刻", "九州エリアプライス平均（円/kWh）")
ws.add_chart(ch, "G4")

# ============================================================
# f2: 0.01円コマ数 / f3: TB4hスプレッド（年度集計を共用）
# ============================================================
fy_agg = daily[daily["fy"].isin(FY_FULL)].groupby("fy").agg(
    floor=("floor_koma", "sum"), curt=("curtail_day", "sum"),
    tb4h_med=("spread_tb4h", "median"),
    tb4h_q1=("spread_tb4h", lambda x: x.quantile(.25)),
    tb4h_q3=("spread_tb4h", lambda x: x.quantile(.75)),
    p_mean=("p_mean", "mean"), solar_peak=("solar_peak_mw", "max"),
).reset_index()

f2 = pd.DataFrame({
    "年度": [f"FY{y}" for y in fy_agg["fy"]],
    "0.01円コマ数（年度計）": fy_agg["floor"].astype(int),
    "出力制御実施日数": fy_agg["curt"].astype(int),
})
ws = new_sheet("f2_床コマ数", "f2 下限価格0.01円への張り付きコマ数（30分コマ・年度計）",
               "九州エリアプライスが0.01円/kWhで約定したコマ数。出力制御はFY2018（2018年10月）開始。掲載: 本編P4右。再現: analysis/02")
h0, h1 = write_df(ws, f2)
ch = BarChart()
ch.type = "col"
ch.title = "0.01円/kWh 約定コマ数（年度計）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.series[0].graphicalProperties.solidFill = BLUE
ch.legend = None
style_axes(ch, "年度", "コマ数")
ws.add_chart(ch, "F4")

f3 = pd.DataFrame({
    "年度": [f"FY{y}" for y in fy_agg["fy"]],
    "25%分位": fy_agg["tb4h_q1"], "中央値": fy_agg["tb4h_med"], "75%分位": fy_agg["tb4h_q3"],
})
ws = new_sheet("f3_TB4hスプレッド", "f3 日次Top4h−Bottom4hスプレッドの年度推移（円/kWh）",
               "各日の九州エリアプライス48コマのうち上位8コマ平均−下位8コマ平均。効率・劣化調整前。FY2022は燃料危機。掲載: Appendix A2左。再現: analysis/02")
h0, h1 = write_df(ws, f3, fmts={"25%分位": "0.00", "中央値": "0.00", "75%分位": "0.00"})
ch = LineChart()
ch.title = "TB4hスプレッド 年度分布（中央値と四分位）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, max_col=4, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, ["B7D3F6", "2A78D6", "B7D3F6"])
style_axes(ch, "年度", "円/kWh")
ws.add_chart(ch, "F4")

# ============================================================
# f4: 太陽光普及 × 昼間床張り付き率（散布図）
# ============================================================
noon = j[(j["hour"] >= 10) & (j["hour"] < 14)]
noon_fy = noon[noon["fy"].isin(FY_FULL)].groupby("fy").agg(
    p_noon=("p_kyushu", "mean"),
    floor_share=("p_kyushu", lambda x: (x <= 0.011).mean() * 100),
).reset_index().merge(fy_agg[["fy", "solar_peak"]], on="fy")
f4 = pd.DataFrame({
    "年度": [f"FY{y}" for y in noon_fy["fy"]],
    "太陽光 年度最大実績出力（GW）": noon_fy["solar_peak"] / 1000,
    "昼間(10-14時) 0.01円張り付き率（%）": noon_fy["floor_share"],
    "昼間平均価格（円/kWh）": noon_fy["p_noon"],
})
ws = new_sheet("f4_太陽光×床張り付き", "f4 太陽光普及と昼間価格の崩落",
               "横軸は需給実績の太陽光発電実績の年度最大値（導入量の代理指標）。掲載: 本編P3右。再現: analysis/02")
h0, h1 = write_df(ws, f4, fmts={"太陽光 年度最大実績出力（GW）": "0.00", "昼間(10-14時) 0.01円張り付き率（%）": "0.0", "昼間平均価格（円/kWh）": "0.00"})
ch = ScatterChart()
ch.title = "太陽光普及 × 昼間の0.01円張り付き率"
ch.height, ch.width = 11, 16
ch.scatterStyle = "marker"
xref = Reference(ws, min_col=2, min_row=h0 + 1, max_row=h1)
yref = Reference(ws, min_col=3, min_row=h0, max_row=h1)
s = Series(yref, xref, title_from_data=True)
s.marker = Marker(symbol="circle", size=7)
s.marker.graphicalProperties.solidFill = BLUE
s.marker.graphicalProperties.line.solidFill = BLUE
s.graphicalProperties.line.noFill = True
ch.series.append(s)
style_axes(ch, "太陽光の年度最大実績出力（GW）", "昼間の0.01円張り付き率（%）")
ch.legend = None
ws.add_chart(ch, "G4")

# ============================================================
# f5: 出力制御日 vs 非制御日の平均価格カーブ（FY2023-25）
# ============================================================
dsub = daily[daily["fy"].isin([2023, 2024, 2025])][["day", "curtail_day"]]
jj = j.merge(dsub, left_on=j["ts"].dt.normalize(), right_on="day", how="inner")
n_c = jj[jj["curtail_day"]]["day"].nunique()
n_n = jj[~jj["curtail_day"]]["day"].nunique()
f5 = pd.DataFrame({"時刻": sorted(jj["hour"].unique())})
f5[f"出力制御 実施日（{n_c}日）"] = f5["時刻"].map(jj[jj["curtail_day"]].groupby("hour")["p_kyushu"].mean())
f5[f"非実施日（{n_n}日）"] = f5["時刻"].map(jj[~jj["curtail_day"]].groupby("hour")["p_kyushu"].mean())
ws = new_sheet("f5_制御日vs非制御日", "f5 出力制御日の価格形状（FY2023–25）",
               "出力制御実施日＝需給実績の太陽光・風力抑制量が正の日。掲載: Appendix A2右。再現: analysis/02")
h0, h1 = write_df(ws, f5, fmts={c: "0.00" for c in f5.columns})
ch = LineChart()
ch.title = "時間帯別平均価格：制御日 vs 非制御日（円/kWh）"
ch.height, ch.width = 11, 20
data = Reference(ws, min_col=2, max_col=3, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, [CORAL, BLUE], marker=False)
style_axes(ch, "時刻", "九州エリアプライス平均（円/kWh）")
ws.add_chart(ch, "F4")

# ============================================================
# f6: 価格分位点の年度推移
# ============================================================
qf = j[j["fy"].isin(FY_FULL)].groupby("fy")["p_kyushu"].quantile([.05, .25, .5, .75, .95]).unstack()
f6 = pd.DataFrame({"年度": [f"FY{y}" for y in qf.index]})
for q, lab in [(.05, "5%分位"), (.25, "25%分位"), (.5, "中央値"), (.75, "75%分位"), (.95, "95%分位")]:
    f6[lab] = qf[q].to_numpy()
ws = new_sheet("f6_価格分位点", "f6 価格分布の二極化（30分コマ別価格の年度内分位点）",
               "FY2019以降、5%分位は0.01円に到達。FY2020の95%分位上昇は2021年1月の需給逼迫による。掲載: Appendix A3左。再現: analysis/02")
h0, h1 = write_df(ws, f6, fmts={c: "0.00" for c in f6.columns if c != "年度"})
ch = LineChart()
ch.title = "九州エリアプライスの年度内分位点（円/kWh）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, max_col=6, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, ["CDE2FB", "86B6EF", "184F95", "86B6EF", "CDE2FB"])
style_axes(ch, "年度", "円/kWh")
ws.add_chart(ch, "H4")

# ============================================================
# f7: バックテスト年間裁定粗利（＋捕捉率は数式）
# ============================================================
ann = bt[bt["fy"].isin(FY_FULL)].groupby("fy")[["pf", "naive", "forecast"]].sum() / 1000
f7 = pd.DataFrame({
    "年度": [f"FY{y}" for y in ann.index],
    "完全予見（円/kW-年）": ann["pf"],
    "a. 前日価格ナイーブ（円/kW-年）": ann["naive"],
    "b. 前日予測ベース（円/kW-年）": ann["forecast"],
})
ws = new_sheet("f7_バックテスト", "f7 4時間蓄電池のスポット裁定価値（実績バックテスト）",
               "1MW/4MWh・往復効率85%・日次1サイクル・劣化無視・スポットのみの下限推定。a=前日価格ランク、b=D-2までの実績による残余需要予測。掲載: 本編P5。再現: analysis/03")
h0, h1 = write_df(ws, f7, fmts={c: "#,##0" for c in f7.columns if c != "年度"},
                  width={"完全予見（円/kW-年）": 16, "a. 前日価格ナイーブ（円/kW-年）": 18, "b. 前日予測ベース（円/kW-年）": 16})
ws.cell(row=h0, column=5, value="a捕捉率").font = F_HDR
ws.cell(row=h0, column=5).fill = FILL_HDR
ws.cell(row=h0, column=6, value="b捕捉率").font = F_HDR
ws.cell(row=h0, column=6).fill = FILL_HDR
for i in range(h0 + 1, h1 + 1):
    ws.cell(row=i, column=5, value=f"=C{i}/B{i}").number_format = "0%"
    ws.cell(row=i, column=5).font = F_BODY
    ws.cell(row=i, column=6, value=f"=D{i}/B{i}").number_format = "0%"
    ws.cell(row=i, column=6).font = F_BODY
ch = LineChart()
ch.title = "年間裁定粗利（円/kW-年）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, max_col=4, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, [DBLUE, LBLUE, CORAL])
style_axes(ch, "年度", "円/kW-年")
ws.add_chart(ch, "H4")

# ============================================================
# f7b: 月次平均価格で規格化した年間裁定粗利（追加分析 2026-07-27）
# ============================================================
pm_month = daily.set_index("day")["p_mean"].resample("MS").mean()  # 月次平均価格（日次平均の月平均）
btn = bt.set_index("day").copy()
btn["p_month"] = pm_month.reindex(btn.index, method="ffill")
for c in ["pf", "naive", "forecast"]:
    btn[c + "_n"] = btn[c] / btn["p_month"] / 1000  # (円/MW-日)÷(円/kWh)/1000 → kWh/kW-日
ann_n = btn[btn["fy"].isin(FY_FULL)].groupby("fy")[["pf_n", "naive_n", "forecast_n"]].sum()
f7b = pd.DataFrame({
    "年度": [f"FY{y}" for y in ann_n.index],
    "完全予見（kWh/kW-年）": ann_n["pf_n"],
    "a. 前日価格（kWh/kW-年）": ann_n["naive_n"],
    "b. 前日予測（kWh/kW-年）": ann_n["forecast_n"],
    "（参考）年度平均価格（円/kWh）": daily[daily["fy"].isin(FY_FULL)].groupby("fy")["p_mean"].mean().to_numpy(),
})
ws = new_sheet("f7b_月次規格化粗利", "f7b 月次平均価格で規格化した年間裁定粗利（kWh/kW-年）",
               "各日の裁定粗利を当月の平均価格で除して年度合計した値＝「その月の平均価格の電気の何kWh分を稼いだか」。f10（年度平均で規格化した等価時間）と異なり年度内の価格水準変動（FY2022冬の高騰等）も除去する。追加分析（2026-07-27）。再現: analysis/09")
h0, h1 = write_df(ws, f7b, fmts={c: "#,##0" if "kWh" in c else "0.00" for c in f7b.columns if c != "年度"})
ch = LineChart()
ch.title = "月次規格化した年間裁定粗利（kWh/kW-年）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, max_col=4, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, [DBLUE, LBLUE, CORAL])
style_axes(ch, "年度", "kWh/kW-年")
ws.add_chart(ch, "G4")
print("f7b 月次規格化PF:", ann_n["pf_n"].round(0).to_dict())

# ============================================================
# 揚水運用の変化: 充電の時間帯プロファイルと昼夜逆転（追加分析 2026-07-27）
# ============================================================
pp = panel.copy()
pp["charge"] = -(pp["pumped"].fillna(0).clip(upper=0) + pp["battery"].fillna(0).clip(upper=0))
sel_fy = [2016, 2019, 2022, 2025]
prof = pp[pp["fy"].isin(sel_fy)].groupby(["hour", "fy"])["charge"].mean().unstack()
pw_a = pd.DataFrame({"時刻": prof.index})
for fy in sel_fy:
    pw_a[f"FY{fy}"] = prof[fy].to_numpy()
day_night = pp[pp["fy"].isin(FY_FULL)].copy()
dn = pd.DataFrame({
    "年度": [f"FY{y}" for y in FY_FULL],
    "昼間（10-14時）充電 平均MW": day_night[(day_night["hour"] >= 10) & (day_night["hour"] < 14)].groupby("fy")["charge"].mean().to_numpy(),
    "夜間（1-5時）充電 平均MW": day_night[(day_night["hour"] >= 1) & (day_night["hour"] < 5)].groupby("fy")["charge"].mean().to_numpy(),
})
ws = new_sheet("揚水運用の変化", "揚水（＋蓄電池）充電運用の変化：昼間吸収の面的拡大と夜間汲み上げの縮小",
               "充電＝需給実績の揚水・蓄電池列の負値の絶対値（MW平均）。蓄電池列は2024年3月以降のみ分離、以前は「揚水等」に包含のため合算で集計。発表P8「揚水の面的拡大」の裏付け。追加分析（2026-07-27）。再現: analysis/09")
a0, a1 = write_df(ws, pw_a, fmts={c: "#,##0" for c in pw_a.columns if c != "時刻"})
b0, b1 = write_df(ws, dn, start_col=7, fmts={c: "#,##0" for c in dn.columns if c != "年度"})
ch = LineChart()
ch.title = "時間帯別の充電量プロファイル（年度平均、MW）"
ch.height, ch.width = 10, 15
data = Reference(ws, min_col=2, max_col=1 + len(sel_fy), min_row=a0, max_row=a1)
cats = Reference(ws, min_col=1, min_row=a0 + 1, max_row=a1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, ["B7D3F6", "6DA7EC", "2A78D6", "184F95"], marker=False)
style_axes(ch, "時刻", "充電（汲み上げ）平均MW")
ws.add_chart(ch, "K4")
ch2 = LineChart()
ch2.title = "昼間 vs 夜間の充電量（年度平均、MW）"
ch2.height, ch2.width = 10, 15
data = Reference(ws, min_col=8, max_col=9, min_row=b0, max_row=b1)
cats = Reference(ws, min_col=7, min_row=b0 + 1, max_row=b1)
ch2.add_data(data, titles_from_data=True)
ch2.set_categories(cats)
line_series_colors(ch2, [CORAL, DBLUE])
style_axes(ch2, "年度", "充電 平均MW")
ws.add_chart(ch2, "K25")
print("揚水 昼間充電:", dn.set_index("年度")["昼間（10-14時）充電 平均MW"].round(0).to_dict())
print("揚水 夜間充電:", dn.set_index("年度")["夜間（1-5時）充電 平均MW"].round(0).to_dict())

# ============================================================
# f8: モンテカルロ（再計算・固定シードで公表値を再現）
# ============================================================
rng = np.random.default_rng(20260726)
N_PATHS = 5000
btm = bt.copy()
btm["month"] = btm["day"].dt.month


def fit_and_simulate(df):
    m = df.set_index("day")["pf"]
    month = df.set_index("day")["month"]
    s = m.groupby(month).mean()
    x = m - month.map(s)
    x0, x1 = x.iloc[:-1].to_numpy(), x.iloc[1:].to_numpy()
    rho = float(np.dot(x0, x1) / np.dot(x0, x0))
    eps = x1 - rho * x0
    cal = pd.date_range("2027-04-01", periods=365, freq="D").month
    s_cal = np.array([s[mo] for mo in cal])
    sims = np.empty(N_PATHS)
    for i in range(N_PATHS):
        e = rng.choice(eps, size=365, replace=True)
        xs = np.empty(365)
        prev = 0.0
        for t in range(365):
            prev = rho * prev + e[t]
            xs[t] = prev
        sims[i] = np.maximum(0.0, s_cal + xs).sum() / 1000
    return rho, sims


mc = {}
for label, fys in [("FY2023-25（主）", [2023, 2024, 2025]), ("FY2021-25（感応度）", [2021, 2022, 2023, 2024, 2025])]:
    df = btm[btm["fy"].isin(fys)].dropna(subset=["pf"])
    rho, sims = fit_and_simulate(df)
    mc[label] = (rho, sims)

tab8 = pd.DataFrame({
    "推定窓": list(mc.keys()),
    "AR(1) ρ": [round(v[0], 2) for v in mc.values()],
    "平均（円/kW-年）": [v[1].mean() for v in mc.values()],
    "5%点": [np.percentile(v[1], 5) for v in mc.values()],
    "中央値": [np.percentile(v[1], 50) for v in mc.values()],
    "95%点": [np.percentile(v[1], 95) for v in mc.values()],
})
sims_main = mc["FY2023-25（主）"][1]
hist, edges = np.histogram(sims_main, bins=60)
f8 = pd.DataFrame({
    "ビン中央値（円/kW-年）": (edges[:-1] + edges[1:]) / 2,
    "パス数（/5,000）": hist,
})
real = btm[btm["fy"].isin([2023, 2024, 2025])].groupby("fy")["pf"].sum() / 1000
ws = new_sheet("f8_モンテカルロ", "f8 モンテカルロによる年間価値分布（完全予見ベース、5,000パス）",
               "日次スプレッド＝月別季節成分＋AR(1)＋残差ブートストラップ。負の日は休止(0)。固定シードで再現。掲載: 本編P6右下。再現: analysis/04")
t0, t1 = write_df(ws, tab8, fmts={c: "#,##0" for c in tab8.columns if "円" in c or "点" in c or c == "中央値"})
ws.cell(row=t1 + 2, column=1, value="実現値（検証用・円/kW-年）: " + " / ".join(f"FY{k}: {v:,.0f}" for k, v in real.items())).font = F_NOTE
h0, h1 = write_df(ws, f8, start_row=t1 + 4, fmts={"ビン中央値（円/kW-年）": "#,##0", "パス数（/5,000）": "0"})
ch = BarChart()
ch.type = "col"
ch.title = "年間価値の分布（FY2023-25推定、ヒストグラム）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.series[0].graphicalProperties.solidFill = "B7D3F6"
ch.gapWidth = 10
ch.legend = None
style_axes(ch, "年間裁定粗利（円/kW-年）", "パス数")
ws.add_chart(ch, "D8")

# ============================================================
# f9: 相対スプレッド / f10: 等価時間
# ============================================================
d9 = daily.copy()
d9["rel_spread"] = d9["spread_tb4h"] / d9["p_mean"]
fy9 = d9[d9["fy"].isin(FY_FULL)].groupby("fy").agg(
    rs_q1=("rel_spread", lambda x: x.quantile(.25)),
    rs_med=("rel_spread", "median"),
    rs_q3=("rel_spread", lambda x: x.quantile(.75)),
    abs_med=("spread_tb4h", "median"),
    p_mean=("p_mean", "mean"),
).reset_index()
f9 = pd.DataFrame({
    "年度": [f"FY{y}" for y in fy9["fy"]],
    "25%分位": fy9["rs_q1"], "中央値": fy9["rs_med"], "75%分位": fy9["rs_q3"],
    "（参考）絶対スプレッド中央値（円/kWh）": fy9["abs_med"],
    "（参考）年度平均価格（円/kWh）": fy9["p_mean"],
})
ws = new_sheet("f9_相対スプレッド", "f9 価格水準で規格化したスプレッド（TB4h ÷ 当日平均価格、無次元）",
               "相対スプレッド＝各日のTB4hスプレッド÷当日48コマ平均価格。水準（燃料）効果を除いた形状の尺度。掲載: 本編P6左。再現: analysis/05")
h0, h1 = write_df(ws, f9, fmts={c: "0.00" for c in f9.columns if c != "年度"})
ch = LineChart()
ch.title = "相対スプレッドの年度分布（中央値と四分位）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, max_col=4, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, ["B7D3F6", "2A78D6", "B7D3F6"])
style_axes(ch, "年度", "相対スプレッド（無次元）")
ws.add_chart(ch, "H4")

f10 = pd.DataFrame({
    "年度": [f"FY{y}" for y in ann.index],
    "完全予見（円/kW-年）": ann["pf"],
    "a. 前日価格（円/kW-年）": ann["naive"],
    "b. 前日予測（円/kW-年）": ann["forecast"],
    "年度平均価格（円/kWh）": fy9.set_index("fy").loc[ann.index, "p_mean"].to_numpy(),
})
ws = new_sheet("f10_等価時間", "f10 価格水準で規格化した蓄電池価値（等価時間、kWh/kW-年）",
               "等価時間＝年間裁定粗利÷年度平均価格。「平均価格の電気の何kWh分を稼いだか」。参考: 年間放電量は最大約1,241kWh/kW-年。掲載: Appendix A4左。再現: analysis/05")
h0, h1 = write_df(ws, f10, fmts={c: "#,##0" if "円/kW-年" in c else "0.00" for c in f10.columns if c != "年度"})
for ci, lab in [(6, "完全予見 等価時間"), (7, "a 等価時間"), (8, "b 等価時間")]:
    c = ws.cell(row=h0, column=ci, value=lab + "（kWh/kW-年）")
    c.font = F_HDR
    c.fill = FILL_HDR
    ws.column_dimensions[get_column_letter(ci)].width = 14
for i in range(h0 + 1, h1 + 1):
    for ci, src in [(6, "B"), (7, "C"), (8, "D")]:
        ws.cell(row=i, column=ci, value=f"={src}{i}/E{i}").number_format = "#,##0"
        ws.cell(row=i, column=ci).font = F_BODY
ch = LineChart()
ch.title = "等価時間（年間粗利 ÷ 年度平均価格、kWh/kW-年）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=6, max_col=8, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, [DBLUE, LBLUE, CORAL])
style_axes(ch, "年度", "kWh/kW-年")
ws.add_chart(ch, "J4")

# ============================================================
# f11: 昼間(10-14時)の需給バランス
# ============================================================
panel["curt_total"] = panel["solar_curt"].fillna(0) + panel["wind_curt"].fillna(0)
noonp = panel[(panel["hour"] >= 10) & (panel["hour"] < 14) & panel["fy"].isin(FY_FULL)].copy()
noonp["storage_charge"] = -(noonp["pumped"].fillna(0).clip(upper=0) + noonp["battery"].fillna(0).clip(upper=0))
noonp["export"] = -noonp["interconn"].clip(upper=0)
abs_tab = noonp.groupby("fy").agg(
    storage_charge=("storage_charge", "mean"), export=("export", "mean"),
    curt=("curt_total", "mean"), demand=("demand", "mean"), solar=("solar", "mean"),
)
f11 = pd.DataFrame({
    "年度": [f"FY{y}" for y in abs_tab.index],
    "揚水+蓄電池 充電（GW）": abs_tab["storage_charge"] / 1000,
    "域外送電・関門（GW）": abs_tab["export"] / 1000,
    "出力制御（GW）": abs_tab["curt"] / 1000,
    "エリア需要（GW）": abs_tab["demand"] / 1000,
    "太陽光実績・制御後（GW）": abs_tab["solar"] / 1000,
})
ws = new_sheet("f11_昼間需給", "f11 昼間（10-14時）の需給バランスの変化（MW平均をGW表示）",
               "需給実績より。揚水+蓄電池充電＝負値の絶対値（蓄電池列は2024年3月以降のみ分離）。域外送電＝連系線流出。掲載: 本編P8左。再現: analysis/06")
h0, h1 = write_df(ws, f11, fmts={c: "0.00" for c in f11.columns if c != "年度"})
ch = LineChart()
ch.title = "余剰の吸収先（昼間平均、GW）"
ch.height, ch.width = 10, 15
data = Reference(ws, min_col=2, max_col=4, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, [DBLUE, LBLUE, CORAL])
style_axes(ch, "年度", "GW")
ws.add_chart(ch, "H4")
ch2 = LineChart()
ch2.title = "需要と太陽光（昼間平均、GW）"
ch2.height, ch2.width = 10, 15
data = Reference(ws, min_col=5, max_col=6, min_row=h0, max_row=h1)
ch2.add_data(data, titles_from_data=True)
ch2.set_categories(cats)
line_series_colors(ch2, [DBLUE, CORAL])
style_axes(ch2, "年度", "GW")
ws.add_chart(ch2, "H25")

# ============================================================
# f12: Top4h / Bottom4h 年度中央値
# ============================================================
j["day"] = j["ts"].dt.normalize()


def tb(g):
    p = g.sort_values().to_numpy()
    return pd.Series({"bot4h": p[:8].mean(), "top4h": p[-8:].mean()})


tbd = j[j["fy"].isin(FY_FULL)].groupby(["fy", "day"])["p_kyushu"].apply(lambda g: tb(g)).unstack()
tby = tbd.groupby("fy").median()
f12 = pd.DataFrame({
    "年度": [f"FY{y}" for y in tby.index],
    "Top4h（夕方ピーク側）": tby["top4h"],
    "Bottom4h（昼間の底側）": tby["bot4h"],
})
ws = new_sheet("f12_TopBottom4h", "f12 スプレッドの分解：日次Top4h／Bottom4h平均価格の年度中央値（円/kWh）",
               "各日の上位8コマ平均（Top4h）と下位8コマ平均（Bottom4h）の年度中央値。掲載: 本編P7左。再現: analysis/06")
h0, h1 = write_df(ws, f12, fmts={c: "0.00" for c in f12.columns if c != "年度"})
ch = LineChart()
ch.title = "Top4h / Bottom4h の年度中央値（円/kWh）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, max_col=3, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, [CORAL, BLUE])
style_axes(ch, "年度", "円/kWh")
ws.add_chart(ch, "F4")

# ============================================================
# f13: 昼間価格帯構成（積み上げ）
# ============================================================
noon13 = j[(j["hour"] >= 10) & (j["hour"] < 14) & j["fy"].isin(FY_FULL)]
bands13 = pd.DataFrame({
    "≤0.01円": noon13.groupby("fy")["p_kyushu"].apply(lambda x: (x <= 0.011).mean()),
    "0.01–1円": noon13.groupby("fy")["p_kyushu"].apply(lambda x: ((x > 0.011) & (x <= 1)).mean()),
    "1–3円": noon13.groupby("fy")["p_kyushu"].apply(lambda x: ((x > 1) & (x <= 3)).mean()),
    "3–5円": noon13.groupby("fy")["p_kyushu"].apply(lambda x: ((x > 3) & (x <= 5)).mean()),
    ">5円": noon13.groupby("fy")["p_kyushu"].apply(lambda x: (x > 5).mean()),
}) * 100
f13 = pd.DataFrame({"年度": [f"FY{y}" for y in bands13.index]})
for c in bands13.columns:
    f13[c + "（%）"] = bands13[c].to_numpy()
ws = new_sheet("f13_価格帯構成", "f13 昼間（10-14時）コマの価格帯構成（%）",
               "床コマの減少分は低価格帯に残らず5円超帯へ移った（＝余剰時間そのものが減少）。掲載: Appendix A3右。再現: analysis/06")
h0, h1 = write_df(ws, f13, fmts={c: "0.0" for c in f13.columns if c != "年度"})
ch = BarChart()
ch.type = "col"
ch.grouping = "stacked"
ch.overlap = 100
ch.title = "昼間コマの価格帯構成（%・積み上げ）"
ch.height, ch.width = 11, 18
data = Reference(ws, min_col=2, max_col=6, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
for s, col in zip(ch.series, RAMP5):
    s.graphicalProperties.solidFill = col
    s.graphicalProperties.line.noFill = True
style_axes(ch, "年度", "構成比（%）")
ws.add_chart(ch, "H4")

# ============================================================
# f14: 反実仮想（07のロジックを再現）
# ============================================================
p14 = panel.copy()
p14["day"] = p14["ts"].dt.normalize()
p14["month"] = p14["ts"].dt.month
p14["season"] = ((p14["month"] % 12) // 3)
night = p14[(p14["hour"] >= 2) & (p14["hour"] < 5)]
dn = night.groupby("fy")["demand"].mean()
nuc = p14.groupby("fy")["nuclear"].mean()
deltas = {fy: {"dD": dn[fy] - dn[2023], "dN": nuc[2023] - nuc[fy]} for fy in [2024, 2025]}
est = p14[p14["fy"].isin([2024, 2025])].dropna(subset=["p_kyushu"]).copy()
est["net"] = est["demand"] - est["solar"] - est["wind"] - est["nuclear"]


def fit_g(df, nbins=80):
    q = np.quantile(df["net"], np.linspace(0, 1, nbins + 1))
    q = np.unique(q)
    idx = np.clip(np.searchsorted(q, df["net"], side="right") - 1, 0, len(q) - 2)
    med = df.groupby(idx)["p_kyushu"].median()
    x = (q[:-1] + q[1:]) / 2
    x = x[med.index]
    y = np.maximum.accumulate(med.to_numpy())
    return x, y


G14 = {s: fit_g(est[est["season"] == s]) for s in range(4)}


def apply_g(net, season):
    x, y = G14[season]
    return max(float(np.interp(net, x, y, left=0.01, right=y[-1])), 0.01)


tgt = est.copy()
tgt["p_fit"] = [apply_g(n, s) for n, s in zip(tgt["net"], tgt["season"])]
tgt["dD"] = tgt["fy"].map(lambda f: deltas[f]["dD"])
tgt["dN"] = tgt["fy"].map(lambda f: deltas[f]["dN"])
tgt["net_cf"] = tgt["net"] - tgt["dD"] - tgt["dN"]
tgt["p_cf"] = [apply_g(n, s) for n, s in zip(tgt["net_cf"], tgt["season"])]

d25 = tgt[tgt["fy"] == 2025]
f14a = pd.DataFrame({"時刻": range(24)})
f14a["実績（円/kWh）"] = f14a["時刻"].map(d25.groupby("hour")["p_kyushu"].mean())
f14a["反実仮想（円/kWh）"] = f14a["時刻"].map(d25.groupby("hour")["p_cf"].mean())
floor_counts = {lab: tgt.groupby("fy")[col].apply(lambda x: int((x <= 0.011).sum()))
                for lab, col in [("実績", "p_kyushu"), ("モデル再現", "p_fit"), ("反実仮想", "p_cf")]}
f14b = pd.DataFrame({"年度": ["FY2024", "FY2025"]} | {lab: [v[2024], v[2025]] for lab, v in floor_counts.items()})

ws = new_sheet("f14_反実仮想", "f14 反実仮想：需要増なし・原子力FY2023水準ならFY2024-25の床はもっと深かった",
               f"経験的供給曲線（net=需要−太陽光−風力−原子力）上の部分均衡評価（床は上限方向）。δD（夜間需要増分）: FY2024 {deltas[2024]['dD']:.0f}MW / FY2025 {deltas[2025]['dD']:.0f}MW、δN（原子力減少分）: FY2024 {deltas[2024]['dN']:.0f}MW / FY2025 {deltas[2025]['dN']:.0f}MW。掲載: 本編P9下。再現: analysis/07")
a0, a1 = write_df(ws, f14a, fmts={c: "0.00" for c in f14a.columns if c != "時刻"})
b0, b1 = write_df(ws, f14b, start_col=5)
ch = LineChart()
ch.title = "FY2025の時間帯別平均価格：実績 vs 反実仮想"
ch.height, ch.width = 10, 15
data = Reference(ws, min_col=2, max_col=3, min_row=a0, max_row=a1)
cats = Reference(ws, min_col=1, min_row=a0 + 1, max_row=a1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
line_series_colors(ch, [BLUE, CORAL], marker=False)
style_axes(ch, "時刻", "円/kWh")
ws.add_chart(ch, "J4")
ch2 = BarChart()
ch2.type = "col"
ch2.title = "床（≤0.01円）時間数：実績 vs 反実仮想"
ch2.height, ch2.width = 10, 15
data = Reference(ws, min_col=6, max_col=8, min_row=b0, max_row=b1)
cats = Reference(ws, min_col=5, min_row=b0 + 1, max_row=b1)
ch2.add_data(data, titles_from_data=True)
ch2.set_categories(cats)
for s, col in zip(ch2.series, [DBLUE, LBLUE, CORAL]):
    s.graphicalProperties.solidFill = col
style_axes(ch2, "年度", "時間数")
ws.add_chart(ch2, "J25")

# ============================================================
# f15: 風力の帯域分解（＋2025年1月の実波形）
# ============================================================
p15 = panel[(panel["ts"] >= "2022-04-01") & (panel["ts"] < "2026-04-01")].set_index("ts")


def band_shares(x):
    x = x.interpolate()
    ma6 = x.rolling(6, center=True, min_periods=3).mean()
    ma24 = x.rolling(24, center=True, min_periods=12).mean()
    ma168 = x.rolling(168, center=True, min_periods=84).mean()
    comps = {"<6時間": x - ma6, "6–24時間": ma6 - ma24,
             "1–7日": ma24 - ma168, ">7日": ma168 - x.mean()}
    tot = x.var()
    return {k: float(v.var() / tot * 100) for k, v in comps.items()}


w_sh, s_sh = band_shares(p15["wind"]), band_shares(p15["solar"])
f15b = pd.DataFrame({
    "帯域": list(w_sh.keys()),
    "風力 分散シェア（%）": list(w_sh.values()),
    "太陽光 分散シェア（%）": [s_sh[k] for k in w_sh.keys()],
})
smp = p15.loc["2025-01-01":"2025-01-31", "wind"]
ma24s = smp.rolling(24, center=True, min_periods=12).mean()
f15a = pd.DataFrame({
    "日時": smp.index.strftime("%m/%d %H:%M"),
    "風力出力（MW）": smp.values,
    "24時間移動平均（MW）": ma24s.values,
})
ws = new_sheet("f15_風力帯域分解", "f15 フリート集約後の風力変動は「数日帯域」が支配的（九州 FY2022–25）",
               "移動平均カスケードによる帯域分解（<6h＝x−MA6h等）。非直交分解のためシェア合計は100%にならない（交差項約2割）。掲載: 本編P14。再現: analysis/08")
b0, b1 = write_df(ws, f15b, fmts={c: "0.0" for c in f15b.columns if c != "帯域"})
ch = BarChart()
ch.type = "col"
ch.title = "変動エネルギーの時間スケール配分（分散シェア%）"
ch.height, ch.width = 10, 14
data = Reference(ws, min_col=2, max_col=3, min_row=b0, max_row=b1)
cats = Reference(ws, min_col=1, min_row=b0 + 1, max_row=b1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
for s, col in zip(ch.series, [BLUE, CORAL]):
    s.graphicalProperties.solidFill = col
style_axes(ch, "帯域", "分散シェア（%）")
ws.add_chart(ch, "E4")
a0, a1 = write_df(ws, f15a, start_row=b1 + 3, fmts={"風力出力（MW）": "#,##0", "24時間移動平均（MW）": "#,##0"})
ch2 = LineChart()
ch2.title = "集約出力の実波形（2025年1月・毎時）"
ch2.height, ch2.width = 10, 22
data = Reference(ws, min_col=2, max_col=3, min_row=a0, max_row=a1)
cats = Reference(ws, min_col=1, min_row=a0 + 1, max_row=a1)
ch2.add_data(data, titles_from_data=True)
ch2.set_categories(cats)
line_series_colors(ch2, [BLUE, CORAL], marker=False)
ch2.series[0].graphicalProperties.line.width = 10000
style_axes(ch2, "日時", "九州エリア 風力出力（MW）")
ws.add_chart(ch2, "E22")

# ============================================================
# fig1: 風力導入比率 × 域外融通能力（散布図）
# ============================================================
fig1_data = [
    ("北海道", 21.3, 32.2, ""),
    ("東北", 51.2, 18.6, ""),
    ("東京", 14.4, 0.5, "†"),
    ("中部", 20.9, 1.6, "†"),
    ("関西", 28.5, 0.6, "†"),
    ("中国", 54.7, 3.4, "†"),
    ("四国", 56.0, 6.0, "†"),
    ("九州", 17.6, 3.2, "†"),
]
fg1 = pd.DataFrame({
    "エリア": [n + f for n, _, _, f in fig1_data],
    "連系線容量÷最大需要（%）": [x for _, x, _, _ in fig1_data],
    "風力接続量÷最大需要（%）": [y for _, _, y, _ in fig1_data],
})
ws = new_sheet("fig1_風力×連系線", "図1 風力導入比率 × 域外融通能力：北海道だけが「風力大×融通小」に位置する",
               "†は暫定値（連系線容量に2017年度公表値等を含む）。風力接続量: 系統WG資料（2024年9月末〜2025年9月末、九州は推定）。最大需要: OCCTO需要想定2025年度推定実績（送電端）。連系線: OCCTO運用容量（隣接エリア主要境界の合計）。一次資料での再計算前の試算。掲載: 本編P12左。出典整理: wind-led-storage-evidence.md")
h0, h1 = write_df(ws, fg1, fmts={c: "0.0" for c in fg1.columns if c != "エリア"},
                  width={"エリア": 10, "連系線容量÷最大需要（%）": 22, "風力接続量÷最大需要（%）": 22})
ch = ScatterChart()
ch.title = "風力導入比率 × 域外融通能力"
ch.height, ch.width = 12, 16
ch.scatterStyle = "marker"
# 北海道（1行目）とその他でシリーズを分ける
xref_h = Reference(ws, min_col=2, min_row=h0 + 1, max_row=h0 + 1)
yref_h = Reference(ws, min_col=3, min_row=h0 + 1, max_row=h0 + 1)
s1 = Series(yref_h, xref_h, title="北海道")
s1.marker = Marker(symbol="circle", size=10)
s1.marker.graphicalProperties.solidFill = CORAL
s1.marker.graphicalProperties.line.solidFill = CORAL
s1.graphicalProperties.line.noFill = True
ch.series.append(s1)
xref_o = Reference(ws, min_col=2, min_row=h0 + 2, max_row=h1)
yref_o = Reference(ws, min_col=3, min_row=h0 + 2, max_row=h1)
s2 = Series(yref_o, xref_o, title="その他エリア")
s2.marker = Marker(symbol="circle", size=7)
s2.marker.graphicalProperties.solidFill = BLUE
s2.marker.graphicalProperties.line.solidFill = BLUE
s2.graphicalProperties.line.noFill = True
ch.series.append(s2)
style_axes(ch, "連系線容量 ÷ 最大需要（%）→大きいほど域外融通が容易", "風力接続量 ÷ 最大需要（%）")
ws.add_chart(ch, "F4")

# ============================================================
# 先行研究マップ / 結果×先行研究（related-work-review.md の表を転載）
# ============================================================
def text_table(name, title, note, headers, rows, widths):
    ws2 = new_sheet(name, title, note)
    r0 = 4
    for ci, h in enumerate(headers):
        c = ws2.cell(row=r0, column=1 + ci, value=h)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws2.column_dimensions[get_column_letter(1 + ci)].width = widths[ci]
    for ri, row in enumerate(rows):
        maxlines = 1
        for ci, v in enumerate(row):
            c = ws2.cell(row=r0 + 1 + ri, column=1 + ci, value=v)
            c.font = F_BODY
            c.alignment = Alignment(vertical="top", wrap_text=True)
            maxlines = max(maxlines, -(-len(str(v)) // max(4, int(widths[ci] / 2))))
        ws2.row_dimensions[r0 + 1 + ri].height = max(16, maxlines * 14 + 4)
    return ws2


text_table(
    "先行研究マップ",
    "先行研究マップ — 6つの系譜（詳細・書誌は related-work-review.md / notes 03・04・09・11・12）",
    "A〜C=再エネ→価格、D=蓄電池→価格・均衡、E〜F=測定手段。本分析との関係の詳細は「結果×先行研究」シート",
    ["系譜", "文献", "市場・対象", "本研究に関係する主要知見", "本分析との関係（参照）"],
    [
        ["A. メリットオーダー効果", "Sensfuß et al. (2008, Energy Policy)", "ドイツ", "再エネ優先給電の価格押し下げを初めて定量化（MOE研究の原点）", "f1・f2の背景（昼間の価格低下）"],
        ["A", "Cludius et al. (2014, Energy Econ)", "ドイツ", "MOE 約6→10€/MWh（2010→12年）と拡大を推計", "同上"],
        ["A", "Woo et al. (2011, Energy Policy)", "ERCOT・風力5-8%", "風力は価格水準を下げ分散を拡大（トレードオフ）", "P13浸透率アンカー：北海道の現在シェアが検出水準"],
        ["B. ボラティリティ非対称性", "Ketterer (2014, Energy Econ)", "ドイツ・風力6-8%", "ARX-GARCHで風力→日次ボラ増", "P13アンカー・北海道RQの根拠"],
        ["B", "Rintamäki et al. (2017, Energy Econ)", "独・デンマーク", "太陽光は日中ボラ減・風力は日次ボラ増 — 電源種で符号が異なる", "f15の太陽光/風力対照の価格版。九州→北海道転換の理論的根拠"],
        ["C. 日本・JEPX", "Maekawa et al. (2018, Energies)", "JEPX", "太陽光の昼間価格押し下げ（日本版ダックカーブ）", "f1の先行確認"],
        ["C", "Sakaguchi & Fujii (2021, Frontiers in Sustainability)", "JEPXエリア別 FY2016-19", "分位点回帰。風力MOEは増大・北海道では高価格分位で最大", "唯一の北海道×風力の直接先行研究 → 本研究が倍増後データで更新（P11・P13）"],
        ["C", "Fuke & Ohashi (2025, J. Commodity Markets)", "九州 2016-2020", "太陽光MOE＋ボラ効果の季節性", "九州分析の直接先行研究（期間をFY2026まで延長・床/裁定価値の角度を追加）"],
        ["C", "Applied Economics (2023, 55(18))", "日本エリア間", "東北再エネが連系線経由で東京の価格・ボラに影響", "北本分断レジーム設計の参照点（P12）"],
        ["C", "Rassi & Kanamura (2023, Energy Policy)", "JEPX入札カーブ", "2021年1月スパイクの構造分析", "f6のFY2020分位点上昇の背景"],
        ["D. 蓄電池・均衡", "Sioshansi et al. (2009, Energy Econ)", "PJM", "裁定価値推定の古典。貯蔵は自らの価値を侵食（カニバリゼーション）", "f7の価値水準の比較点・#6の起点"],
        ["D", "Lamp & Samano (2022, Energy Econ)", "CAISO", "蓄電池導入が日中スプレッドを縮小させたことを実証", "f11・f12：九州では揚水が同型の圧縮を先行（予行演習）"],
        ["D", "Butters, Dorsey & Gowrisankaran (2025, Econometrica)", "CAISO 2015-19", "自由参入均衡モデル。不確実性下の価値≈完全予見の70%・劣化▲27%・補助なしでは普及せず。再現パッケージ公開", "最重要ベンチマーク：f7・f8と整合。北海道仕様への拡張が修論本体（P15）"],
        ["D", "Schmalensee (2022)・Junge et al. (2022, Energy J)", "理論", "競争均衡の貯蔵容量＝社会最適・均衡でゼロ利潤", "均衡量K*の定義（P15の③均衡）"],
        ["D", "Karaduman (2021, MIT CEEPR WP)", "南豪州", "価格インパクト無視は貯蔵収益を約2倍過大評価", "価格テイカー仮定の限界の注記（f7脚注）"],
        ["D", "Rangarajan et al. (2023, Energy Econ)", "豪州NEM", "蓄電池はまず調整力（FCAS）費用を低減", "スポット単独=下限推定の根拠（f7）"],
        ["D", "市場実績: Modo Energy・WattClarity (2025-26)", "CAISO・NEM", "蓄電池収益2023→25年に約半減・スプレッド半減", "カニバリの現実性。日本は拡大局面=事前推定の価値（P12）"],
        ["E. 空間平滑化・時間構造", "St. Martin et al. (2015, ERL)", "豪・加・米", "相関距離は時間スケール依存：高周波のみ集約で消える", "f15の理論的根拠"],
        ["E", "Malvaldi et al. (2017, Wind Energy)", "欧州10カ国", "デンマーク〜独の国全体集約でも相関0.65", "f15・P14：打ち消し合い反論への回答"],
        ["E", "Ohlendorf & Schill (2020, ERL 15:084045)", "ドイツ40年", "低風力イベントは毎年約5日連続・10年に1度8日", "低出力158時間（f15）と整合。duration論点"],
        ["F. 計量手法", "Ciarreta, Muniain & Zarraga (2017, J. Forecasting)", "独墺当日市場", "RV=CV+JV分解＋HARで予測改善。BNS検定が最良", "f8の限界→北海道第1段階（HAR-CV-JV-X）の手順書"],
        ["F", "Haugom et al. (2011, Energy Econ)", "Nord Pool", "HAR-RV/CJの電力市場初適用", "同上"],
        ["F", "Weron (2014, IJF)", "レビュー", "電力価格モデリングの標準レビュー", "手法選択の位置づけ"],
    ],
    [16, 30, 16, 42, 40],
)

text_table(
    "結果×先行研究",
    "本分析の結果 × 先行研究の対応（詳細は related-work-review.md §2）",
    "関係の型: 追試=既知現象の日本/最新データでの確認、整合=先行結論と符合、拡張=新対象・新角度へ、新規=対応する先行研究なし、動機=先行手法の必要性を実証",
    ["#", "本分析の結果（図）", "型", "対応する先行研究", "関連の内容"],
    [
        ["1", "ダックカーブ深化・床コマ急増（f1・f2・f4）", "追試・拡張", "Maekawa 2018／Fuke & Ohashi 2025／系譜A", "MOE・昼間押し下げを10年半・下限0.01円張り付きの形で確認。観察期間をFY2026まで延長"],
        ["2", "床・スプレッドがFY2023ピーク比で反転減少（f2・f3）", "新規", "（対応文献なし）", "海外の反転は蓄電池由来（Lamp & Samano）。九州は蓄電池±10MWの段階で反転しており駆動要因が異なる（→#6）"],
        ["3", "裁定価値約1万円/kW-年＜資本費約2万円（f7）", "整合", "Butters 2025／Sioshansi 2009／Rangarajan 2023", "スポット単独では投資不成立＝CAISOの結論と整合。マルチマーケット評価が本番課題"],
        ["4", "実行可能戦略の捕捉率76〜84%（f7）", "整合", "Butters 2025", "「不確実性下の価値は完全予見の約70%」とほぼ同水準を簡便な予測則で確認"],
        ["5", "戦略bが捕捉率約4割→80-84%で逆転（f7）", "新規", "（Sakaguchi & Fujii 2021の裏付けを兼ねる）", "価格形状の予測可能性レジーム変化を運用成績で示す実証装置。文献に先例なし"],
        ["6", "底上げ型圧縮：揚水が限界帯+553MW（f11・f12）", "拡張", "Lamp & Samano 2022／Sioshansi 2009", "CAISOの「蓄電池によるスプレッド圧縮」を、九州では揚水が先行して起こす＝カニバリゼーションの予行演習（異技術間の先食い）"],
        ["7", "制御時間中の床約定率74%→40%（第3報）", "新規", "（Buttersはcurtailment非モデル化と明記）", "物理（制御）と市場（床）の連動低下という観察。出力制御×価格形成の実証は空白"],
        ["8", "フラット需要増は床+20〜33%でも価値ほぼ中立（f14）", "整合・拡張", "Butters 2025（分配効果）", "「昼だけ効く吸収だけが蓄電池価値を侵食する」機構を部分均衡の会計分解で先取り"],
        ["9", "MCが年度間変動を過小評価（f8）", "動機", "Ciarreta et al. 2017／Haugom et al. 2011", "HAR型長期記憶の導入（北海道第1段階）の実証的動機付け"],
        ["10", "風力帯域分解：<6h 1.9%・1-7日43.7%（f15）", "追試・拡張", "St. Martin 2015／Malvaldi 2017／Ohlendorf & Schill 2020／Rintamäki 2017", "空間平滑化の帯域依存を日本の需給実績で確認。分散配分→duration価値の理論上限という読み替えが本研究の視点"],
        ["11", "北海道：風力/需要32%・連系線の1.5倍（図1）", "拡張", "Sakaguchi & Fujii 2021／Woo 2011・Ketterer 2014／Applied Economics 2023", "設備半分時代の検出を倍増後データで更新。浸透率は検出水準に到達しつつある（P13）"],
    ],
    [4, 34, 11, 34, 45],
)

# ============================================================
# 目次シート（先頭へ）
# ============================================================
toc = wb.create_sheet("目次", 0)
toc["A1"] = "ゼミ発表 図表データ集（九州パイロット分析＋北海道転換）"
toc["A1"].font = Font(name=FONT, size=14, bold=True, color="1F3864")
toc["A2"] = "各シートに図の数値データとグラフを収録。データは公開データ（JEPXスポット・九州電力需給実績）から thesis/analysis/02〜08 と同一ロジックで再計算。"
toc["A2"].font = F_NOTE
toc["A3"] = "生成: thesis/analysis/09_export_figures_xlsx.py（2026年7月26日）。図の定義・限界の一覧は kyushu-methods-data.md を参照。"
toc["A3"].font = F_NOTE
rows = [
    ("シート", "図", "内容", "発表資料での掲載箇所"),
    ("先行研究マップ", "—", "先行研究6系譜の一覧と本分析との関係", "本編P11・P13の裏付け"),
    ("結果×先行研究", "—", "本分析の結果11項目と先行研究の対応（追試/整合/拡張/新規/動機）", "想定問答・論文関連研究章"),
    ("f1_ダックカーブ", "f1", "時間帯別平均価格カーブ（FY2016/19/23/25）", "本編P4 左"),
    ("f2_床コマ数", "f2", "0.01円張り付きコマ数の年度推移", "本編P4 右"),
    ("f3_TB4hスプレッド", "f3", "日次TB4hスプレッドの年度分布", "Appendix A2 左"),
    ("f4_太陽光×床張り付き", "f4", "太陽光普及×昼間床張り付き率（散布図）", "本編P3 右"),
    ("f5_制御日vs非制御日", "f5", "出力制御日vs非制御日の価格カーブ", "Appendix A2 右"),
    ("f6_価格分位点", "f6", "価格分位点の年度推移（分布の二極化）", "Appendix A3 左"),
    ("f7_バックテスト", "f7", "蓄電池裁定価値バックテスト（3戦略）", "本編P5"),
    ("f7b_月次規格化粗利", "追加", "月次平均価格で規格化した年間裁定粗利", "追加分析（2026-07-27）"),
    ("f8_モンテカルロ", "f8", "年間価値のモンテカルロ分布", "本編P6 右下"),
    ("f9_相対スプレッド", "f9", "価格水準で規格化したスプレッド", "本編P6 左"),
    ("f10_等価時間", "f10", "規格化年間価値（等価時間）", "Appendix A4 左"),
    ("f11_昼間需給", "f11", "昼間の需給バランスの変化", "本編P8 左"),
    ("揚水運用の変化", "追加", "揚水充電の時間帯プロファイルと昼夜逆転", "追加分析（2026-07-27）"),
    ("f12_TopBottom4h", "f12", "Top4h/Bottom4hの分解", "本編P7 左"),
    ("f13_価格帯構成", "f13", "昼間コマの価格帯構成（積み上げ）", "Appendix A3 右"),
    ("f14_反実仮想", "f14", "反実仮想（需要増なし・原子力FY2023水準）", "本編P9 下"),
    ("f15_風力帯域分解", "f15", "風力変動の帯域分解＋実波形", "本編P14"),
    ("fig1_風力×連系線", "図1", "風力導入比率×域外融通能力（散布図）", "本編P12 左"),
]
for ri, row in enumerate(rows):
    for ci, v in enumerate(row):
        c = toc.cell(row=5 + ri, column=1 + ci, value=v)
        if ri == 0:
            c.font = F_HDR
            c.fill = FILL_HDR
        else:
            c.font = F_BODY
for col, w in [("A", 24), ("B", 8), ("C", 46), ("D", 22)]:
    toc.column_dimensions[col].width = w

order = ["目次", "先行研究マップ", "結果×先行研究",
         "f1_ダックカーブ", "f2_床コマ数", "f3_TB4hスプレッド", "f4_太陽光×床張り付き",
         "f5_制御日vs非制御日", "f6_価格分位点", "f7_バックテスト", "f7b_月次規格化粗利",
         "f8_モンテカルロ", "f9_相対スプレッド", "f10_等価時間", "f11_昼間需給",
         "揚水運用の変化", "f12_TopBottom4h", "f13_価格帯構成", "f14_反実仮想",
         "f15_風力帯域分解", "fig1_風力×連系線"]
wb._sheets = [wb[n] for n in order]

wb.save(OUT)
print("saved", OUT)
