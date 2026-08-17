#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""進捗報告（2026-08-17）デッキの図表データをExcel化（編集可能なネイティブグラフ付き）

出力: thesis/slides/進捗報告_20260817_図表データ.xlsx
シート: 目次 / 分断方向転換 / capture率 / 帯域分解 / 泊DC_2x2 / 文献構成
"""
import os

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.join(os.path.dirname(__file__), "..")
OUT = os.path.join(ROOT, "slides", "進捗報告_20260817_図表データ.xlsx")

BLUE = "0079C2"
TEAL = "176871"
HDR = PatternFill("solid", fgColor=BLUE)
HDR_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=13, color=TEAL)

wb = Workbook()


def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center")


def base_chart(ch, title, ytitle):
    ch.type = "col"
    ch.style = 10
    ch.title = title
    ch.y_axis.title = ytitle
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.gapWidth = 60
    ch.height = 9
    ch.width = 16
    return ch


# ---------- 目次 ----------
ws = wb.active
ws.title = "目次"
ws["A1"] = "進捗報告（2026-08-17）図表データ"
ws["A1"].font = Font(bold=True, size=14, color=TEAL)
toc = [
    ("分断方向転換", "スライド4: 市場分断の方向転換（コマ構成比）"),
    ("capture率", "スライド5: 実行可能戦略のcapture率（完全予見比）"),
    ("帯域分解", "スライド6: 風力変動の帯域分解（分散シェア）"),
    ("泊DC_2x2", "スライド15: 泊3号×DC需要の2×2分解（価格過程v1）"),
    ("文献構成", "スライド7: コア文献41本のセット構成"),
]
ws["A3"] = "シート"
ws["B3"] = "内容"
style_header(ws, 3, 2)
for i, (a, b) in enumerate(toc, start=4):
    ws.cell(row=i, column=1, value=a)
    ws.cell(row=i, column=2, value=b)
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 64
ws["A10"] = "数値の出所: 北海道時間パネル（11_build_area_panels）・北海道簡易分析（10_hokkaido_quick_xlsx）・価格過程v1（13_price_process_v1）"
ws["A10"].font = Font(size=9, color="595959")

# ---------- 分断方向転換 ----------
ws = wb.create_sheet("分断方向転換")
ws["A1"] = "市場分断の方向転換（北海道エリアプライス vs システムプライス、コマ構成比%）"
ws["A1"].font = TITLE_FONT
rows = [
    ("区分", "FY2016-19", "FY2023-25"),
    ("高値分断（道内＞本州）", 89.5, 60.0),
    ("連系（±0.01円以内）", 7.5, 11.0),
    ("安値分断（道内＜本州）", 3.0, 29.0),
]
for r, row in enumerate(rows, start=3):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)
style_header(ws, 3, 3)
ws.column_dimensions["A"].width = 26
ch = base_chart(BarChart(), "市場分断の方向転換（コマ構成比%）", "%")
data = Reference(ws, min_col=2, max_col=3, min_row=3, max_row=6)
cats = Reference(ws, min_col=1, min_row=4, max_row=6)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ws.add_chart(ch, "E3")
ws["A8"] = "注: 高値87-92%等の年度レンジの代表値。判定はシステムプライス±0.01円"
ws["A8"].font = Font(size=9, color="595959")

# ---------- capture率 ----------
ws = wb.create_sheet("capture率")
ws["A1"] = "実行可能戦略のcapture率（完全予見PF=100%、60分粒度）"
ws["A1"].font = TITLE_FONT
rows = [
    ("戦略", "capture率(%)", "レンジ"),
    ("完全予見PF", 100.0, "—"),
    ("九州 戦略a（価格ランク）", 78.5, "77-80%"),
    ("北海道 戦略a（価格ランク）", 67.5, "66-69%"),
    ("北海道 戦略b（残余需要予測）", 79.0, "77-81%"),
]
for r, row in enumerate(rows, start=3):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)
style_header(ws, 3, 3)
ws.column_dimensions["A"].width = 30
ws.column_dimensions["C"].width = 12
ch = base_chart(BarChart(), "capture率（完全予見PF=100%）", "%")
data = Reference(ws, min_col=2, min_row=3, max_row=7)
cats = Reference(ws, min_col=1, min_row=4, max_row=7)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.legend = None
ws.add_chart(ch, "E3")
ws["A9"] = "注: 棒は各年度レンジの中央値。戦略b′（風力持続性の追加）は改善なし"
ws["A9"].font = Font(size=9, color="595959")

# ---------- 帯域分解 ----------
ws = wb.create_sheet("帯域分解")
ws["A1"] = "風力変動の帯域分解（フリート集約・制御前出力の分散シェア%、FY2024-25）"
ws["A1"].font = TITLE_FONT
rows = [
    ("帯域", "分散シェア(%)"),
    ("日内（<24h）", 25.7),
    ("1〜7日", 35.7),
    ("7日超", 38.6),
]
for r, row in enumerate(rows, start=3):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)
style_header(ws, 3, 2)
ws.column_dimensions["A"].width = 18
ch = base_chart(BarChart(), "風力変動の帯域分解（分散シェア%）", "%")
data = Reference(ws, min_col=2, min_row=3, max_row=6)
cats = Reference(ws, min_col=1, min_row=4, max_row=6)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.legend = None
ws.add_chart(ch, "D3")
ws["A8"] = "注: バンドパス分散分解。太陽光は6-24h帯が中心（対照）。4h蓄電池が主に吸収できるのは日内帯のみ"
ws["A8"].font = Font(size=9, color="595959")

# ---------- 泊DC_2x2 ----------
ws = wb.create_sheet("泊DC_2x2")
ws["A1"] = "泊3号×DC需要の2×2分解（価格過程v1・FY2023-25実パス・仕様B・共通乱数）"
ws["A1"].font = TITLE_FONT
rows = [
    ("シナリオ", "床時間(h/年)", "PF価値(円/kW-年)", "ΔPF(円/kW-年)"),
    ("ベース（泊なし・DCなし）", 117, 8593, 0),
    ("泊3号再稼働のみ（91.2万kW×90%）", 704, 9777, 1184),
    ("DC需要+80万kW（フラット）のみ", 29, 7796, -797),
    ("泊＋DC（両方）", 121, 8582, -11),
]
for r, row in enumerate(rows, start=3):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)
style_header(ws, 3, 4)
ws.column_dimensions["A"].width = 34
for col in "BCD":
    ws.column_dimensions[col].width = 16
ch = base_chart(BarChart(), "PF価値の変化（ベース比、円/kW-年）", "円/kW-年")
data = Reference(ws, min_col=4, min_row=3, max_row=7)
cats = Reference(ws, min_col=1, min_row=4, max_row=7)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.legend = None
ws.add_chart(ch, "F3")
ws["A9"] = "注: 部分均衡（蓄電池フリート応答・p_system変化は未反映）。床時間は60分粒度のモデル値（実績の床コマ473/年とは定義が異なる）"
ws["A9"].font = Font(size=9, color="595959")

# ---------- 文献構成 ----------
ws = wb.create_sheet("文献構成")
ws["A1"] = "コア文献41本のセット構成（DOI検証・PDF収集・抽出表・数値チェック済み）"
ws["A1"].font = TITLE_FONT
rows = [
    ("セット", "本数", "内容"),
    ("A 再エネ→価格・ボラ実証", 11, "Woo 2011・Ketterer 2014・Wozabal 2016 ほか"),
    ("B 蓄電池の経済学・均衡", 13, "Butters 2025・Karaduman・Schmalensee ほか"),
    ("C 日本・JEPX", 9, "Fuke & Ohashi 2025・Sakaguchi & Fujii 2021 ほか"),
    ("D 手法・投資評価", 8, "Corsi 2009・Fanone 2013・Leahy/Grenadier ほか"),
]
for r, row in enumerate(rows, start=3):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)
style_header(ws, 3, 3)
ws.column_dimensions["A"].width = 28
ws.column_dimensions["C"].width = 48
ch = base_chart(BarChart(), "コア文献のセット構成（本）", "本")
data = Reference(ws, min_col=2, min_row=3, max_row=7)
cats = Reference(ws, min_col=1, min_row=4, max_row=7)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.legend = None
ws.add_chart(ch, "E3")

wb.save(OUT)
print("saved", OUT)
