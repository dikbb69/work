#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""北海道の時間帯別価格カーブの形状: 3季節 × 太陽光の大小（FY2023-25）

定義:
  需要期・夏   = 7-8月
  需要期・冬   = 12-2月
  不需要期     = 4-5月・10-11月（春・秋）
  太陽光大/小 = 各季節内で「日次の太陽光発電量（制御前）」の上位1/3 / 下位1/3の日
  価格        = 北海道エリアプライス（60分平均）の時刻別平均

出力:
  thesis/figures/hokkaido/price_curve_shapes_fy2023-25.png（3面小倍数）
  slides/進捗報告_20260817_図表データ.xlsx に「価格カーブ形状」シートを追加（編集可能なグラフ付き）
"""
import os

import numpy as np
import pandas as pd
from matplotlib import font_manager
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

for f in ["/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
          "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"]:
    try:
        font_manager.fontManager.addfont(f)
    except Exception:
        pass
plt.rcParams["font.family"] = "Noto Sans CJK JP"

ROOT = os.path.join(os.path.dirname(__file__), "..")
PROC = os.path.join(ROOT, "data", "processed")
FIGDIR = os.path.join(ROOT, "figures", "hokkaido")
XLSX = os.path.join(ROOT, "slides", "進捗報告_20260817_図表データ.xlsx")
os.makedirs(FIGDIR, exist_ok=True)

C_SUN = "#D95B20"   # 太陽光大（検証済みパレット）
C_LOW = "#0079C2"   # 太陽光小
GRAY = "#595959"

p = pd.read_csv(os.path.join(PROC, "hokkaido_hourly_panel.csv"), parse_dates=["ts"])
p = p[(p["ts"] >= "2023-04-01") & (p["ts"] < "2026-04-01")].dropna(subset=["p_hokkaido"]).copy()
p["day"] = p["ts"].dt.normalize()
p["hour"] = p["ts"].dt.hour
p["solar_pre"] = p["solar"].fillna(0) + p["solar_curt"].fillna(0)

SEASONS = [
    ("需要期・夏（7-8月）", [7, 8]),
    ("需要期・冬（12-2月）", [12, 1, 2]),
    ("不需要期（4-5・10-11月）", [4, 5, 10, 11]),
]

# 24時間そろった日のみ
cnt = p.groupby("day")["hour"].count()
p = p[p["day"].isin(cnt[cnt == 24].index)]

curves = {}   # season -> DataFrame(hour, 太陽光大, 太陽光小)
counts = {}
for name, months in SEASONS:
    sub = p[p["ts"].dt.month.isin(months)]
    daily_sun = sub.groupby("day")["solar_pre"].sum()
    q_lo, q_hi = daily_sun.quantile([1 / 3, 2 / 3])
    d_hi = daily_sun[daily_sun >= q_hi].index
    d_lo = daily_sun[daily_sun <= q_lo].index
    hi = sub[sub["day"].isin(d_hi)].groupby("hour")["p_hokkaido"].mean()
    lo = sub[sub["day"].isin(d_lo)].groupby("hour")["p_hokkaido"].mean()
    curves[name] = pd.DataFrame({"太陽光大": hi, "太陽光小": lo})
    counts[name] = (len(d_hi), len(d_lo))
    print(f"{name}: 太陽光大 {len(d_hi)}日 / 小 {len(d_lo)}日 | 昼(11-13時)差 "
          f"{(hi.loc[11:13].mean() - lo.loc[11:13].mean()):+.2f}円/kWh")

# ---------- PNG（3面小倍数・共有y軸） ----------
fig, axes = plt.subplots(1, 3, figsize=(13.5, 4.6), sharey=True)
for ax, (name, _) in zip(axes, SEASONS):
    cv = curves[name]
    ax.plot(cv.index, cv["太陽光大"], color=C_SUN, lw=2.2, label="太陽光大の日（上位1/3）")
    ax.plot(cv.index, cv["太陽光小"], color=C_LOW, lw=2.2, label="太陽光小の日（下位1/3）")
    ax.set_title(name, fontsize=12, pad=8)
    ax.set_xticks([0, 6, 12, 18, 23])
    ax.set_xlabel("時刻", fontsize=10, color=GRAY)
    ax.grid(axis="y", color="#DDDDDD", lw=0.6)
    for sp in ["top", "right"]:
        ax.spines[sp].set_visible(False)
    ax.tick_params(colors=GRAY, labelsize=9.5)
axes[0].set_ylabel("北海道エリアプライス（円/kWh）", fontsize=10, color=GRAY)
axes[0].legend(loc="upper left", fontsize=9.5, frameon=False)
fig.suptitle("北海道の時間帯別価格カーブ — 3季節 × 太陽光の大小（FY2023-25平均）",
             fontsize=13.5, fontweight="bold", y=1.02)
fig.text(0.995, -0.04,
         "太陽光大/小＝各季節内で日次太陽光発電量（出力制御前）の上位/下位1/3の日。60分平均価格の時刻別平均。データ: JEPX・北海道電力NW需給実績",
         ha="right", fontsize=8, color=GRAY)
fig.tight_layout()
png = os.path.join(FIGDIR, "price_curve_shapes_fy2023-25.png")
fig.savefig(png, dpi=160, bbox_inches="tight")
print("saved", png)

# ---------- Excel（編集可能グラフ） ----------
wb = load_workbook(XLSX)
if "価格カーブ形状" in wb.sheetnames:
    del wb["価格カーブ形状"]
ws = wb.create_sheet("価格カーブ形状")
ws["A1"] = "北海道の時間帯別価格カーブ — 3季節 × 太陽光の大小（FY2023-25、円/kWh）"
ws["A1"].font = Font(bold=True, size=13, color="176871")

HDR = PatternFill("solid", fgColor="0079C2")
col0 = 1
chart_anchors = ["A31", "H31", "O31"]
for si, (name, _) in enumerate(SEASONS):
    cv = curves[name]
    c = col0 + si * 3
    ws.cell(row=3, column=c, value=f"{name}")
    ws.cell(row=3, column=c).font = Font(bold=True, size=11)
    heads = ["時刻", "太陽光大", "太陽光小"]
    for k, h in enumerate(heads):
        cell = ws.cell(row=4, column=c + k, value=h)
        cell.fill = HDR
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for r, hour in enumerate(cv.index, start=5):
        ws.cell(row=r, column=c, value=int(hour))
        ws.cell(row=r, column=c + 1, value=round(float(cv.loc[hour, "太陽光大"]), 2))
        ws.cell(row=r, column=c + 2, value=round(float(cv.loc[hour, "太陽光小"]), 2))
    ch = LineChart()
    ch.title = name
    ch.style = 10
    ch.height, ch.width = 8.5, 11.5
    data = Reference(ws, min_col=c + 1, max_col=c + 2, min_row=4, max_row=28)
    cats = Reference(ws, min_col=c, min_row=5, max_row=28)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.x_axis.title = "時刻"
    ch.y_axis.title = "円/kWh"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    for ln, color in zip(ch.series, ["D95B20", "0079C2"]):
        ln.graphicalProperties.line.solidFill = color
        ln.graphicalProperties.line.width = 22000
        ln.smooth = False
    ws.add_chart(ch, chart_anchors[si])

ws["A29"] = ("注: 太陽光大/小＝各季節内で日次太陽光発電量（出力制御前）の上位/下位1/3の日（日数: "
             + " / ".join(f"{n.split('（')[0]} {a}日:{b}日" for n, (a, b) in counts.items())
             + "）。60分平均価格の時刻別平均。季節定義は月で編集可能（analysis/15）")
ws["A29"].font = Font(size=9, color="595959")
for col in range(1, 10):
    ws.column_dimensions[chr(64 + col)].width = 11

# 目次に追記
toc = wb["目次"]
row = toc.max_row + 1
toc.cell(row=row, column=1, value="価格カーブ形状")
toc.cell(row=row, column=2, value="追加: 3季節×太陽光大小の時間帯別価格カーブ（FY2023-25）")
wb.save(XLSX)
print("saved sheet 価格カーブ形状 →", XLSX)
