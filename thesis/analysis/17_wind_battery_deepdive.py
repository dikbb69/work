#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""風力×蓄電池の初期分析（北海道、FY2023-25）

  1. 風力大/小の日の時間帯別価格カーブ（3季節、15の太陽光版の風力版）
  2. 風力の日次水準 × 裁定指標（TB4hスプレッド・PF粗利・床時間・安値分断時間）
  3. duration曲線: 完全予見裁定価値のh依存（2/4/6/8時間、実績価格）
  4. 風力の日内プロファイル（季節別・時刻別平均cf）と風力抑制の実績

出力: figures/hokkaido/*.png ＋ 進捗報告_20260817_図表データ.xlsx に3シート追加
"""
import os

import numpy as np
import pandas as pd
from matplotlib import font_manager
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

for f in ["/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
          "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"]:
    try:
        font_manager.fontManager.addfont(f)
    except Exception:
        pass
plt.rcParams["font.family"] = "Noto Sans CJK JP"

ROOT = os.path.join(os.path.dirname(__file__), "..")
PROC = os.path.join(ROOT, "data", "processed")
FIGDIR = os.path.join(ROOT, "figures", "hokkaido")
XLSX = os.path.join(ROOT, "slides", "進捗報告_20260817_図表データ.xlsx")

C_HI = "#D95B20"   # 風力大（資源が多い日=オレンジ、太陽光図と同じ意味論）
C_LO = "#0079C2"   # 風力小
GRAY = "#595959"
ETA = 0.85

p = pd.read_csv(os.path.join(PROC, "hokkaido_hourly_panel.csv"), parse_dates=["ts"])
p["fy"] = p["ts"].dt.year - (p["ts"].dt.month < 4).astype(int)
p = p[(p["ts"] >= "2023-04-01") & (p["ts"] < "2026-04-01")].dropna(subset=["p_hokkaido"]).copy()
p["day"] = p["ts"].dt.normalize()
p["hour"] = p["ts"].dt.hour
p["wind_pre"] = p["wind"].fillna(0) + p["wind_curt"].fillna(0)
cnt = p.groupby("day")["hour"].count()
p = p[p["day"].isin(cnt[cnt == 24].index)]

SEASONS = [
    ("需要期・夏（7-8月）", [7, 8]),
    ("需要期・冬（12-2月）", [12, 1, 2]),
    ("不需要期（4-5・10-11月）", [4, 5, 10, 11]),
]

# ============================================================
# 1. 風力大/小の日の価格カーブ（3季節）
# ============================================================
curves, counts = {}, {}
for name, months in SEASONS:
    sub = p[p["ts"].dt.month.isin(months)]
    dw = sub.groupby("day")["wind_pre"].sum()
    q_lo, q_hi = dw.quantile([1 / 3, 2 / 3])
    hi = sub[sub["day"].isin(dw[dw >= q_hi].index)].groupby("hour")["p_hokkaido"].mean()
    lo = sub[sub["day"].isin(dw[dw <= q_lo].index)].groupby("hour")["p_hokkaido"].mean()
    curves[name] = pd.DataFrame({"風力大": hi, "風力小": lo})
    counts[name] = (int((dw >= q_hi).sum()), int((dw <= q_lo).sum()))
    print(f"[curve] {name}: 全時間平均差 {(hi.mean() - lo.mean()):+.2f}円/kWh")

fig, axes = plt.subplots(1, 3, figsize=(13.5, 4.6), sharey=True)
for ax, (name, _) in zip(axes, SEASONS):
    cv = curves[name]
    ax.plot(cv.index, cv["風力大"], color=C_HI, lw=2.2, label="風力大の日（上位1/3）")
    ax.plot(cv.index, cv["風力小"], color=C_LO, lw=2.2, label="風力小の日（下位1/3）")
    ax.set_title(name, fontsize=12, pad=8)
    ax.set_xticks([0, 6, 12, 18, 23])
    ax.set_xlabel("時刻", fontsize=10, color=GRAY)
    ax.grid(axis="y", color="#DDDDDD", lw=0.6)
    for sp in ["top", "right"]:
        ax.spines[sp].set_visible(False)
    ax.tick_params(colors=GRAY, labelsize=9.5)
axes[0].set_ylabel("北海道エリアプライス（円/kWh）", fontsize=10, color=GRAY)
axes[0].legend(loc="lower left", fontsize=9.5, frameon=False)
fig.suptitle("風力大の日は一日全体が沈む — 形はほぼ平行移動（太陽光と対照的）（FY2023-25平均）",
             fontsize=13.5, fontweight="bold", y=1.02)
fig.text(0.995, -0.04,
         "風力大/小＝各季節内で日次風力発電量（出力制御前）の上位/下位1/3の日。60分平均価格の時刻別平均。データ: JEPX・北海道電力NW需給実績",
         ha="right", fontsize=8, color=GRAY)
fig.tight_layout()
fig.savefig(os.path.join(FIGDIR, "wind_price_curves_fy2023-25.png"), dpi=160, bbox_inches="tight")
print("saved wind_price_curves")

# ============================================================
# 2. 風力の日次水準 × 裁定指標（季節別・風力三分位）
# ============================================================
P = p.pivot(index="day", columns="hour", values="p_hokkaido")
daily = pd.DataFrame(index=P.index)
srt = np.sort(P.to_numpy(), axis=1)
daily["TB4hスプレッド"] = srt[:, -4:].mean(axis=1) - srt[:, :4].mean(axis=1)
daily["PF粗利"] = np.maximum(0.0, (srt[:, -4:].sum(axis=1) * ETA - srt[:, :4].sum(axis=1)) * 1000)
daily["床時間"] = (P <= 0.011).sum(axis=1)
split = p.assign(low_split=(p["p_hokkaido"] < p["p_system"] - 0.01))
daily["安値分断時間"] = split.groupby("day")["low_split"].sum()
daily["wind_day"] = p.groupby("day")["wind_pre"].sum()
daily["month"] = daily.index.month

rows2 = []
for name, months in SEASONS:
    d = daily[daily["month"].isin(months)].copy()
    q = d["wind_day"].quantile([1 / 3, 2 / 3])
    d["ter"] = np.select([d["wind_day"] <= q.iloc[0], d["wind_day"] >= q.iloc[1]], ["風力小", "風力大"], "中位")
    g = d.groupby("ter")[["TB4hスプレッド", "PF粗利", "床時間", "安値分断時間"]].mean()
    for ter in ["風力小", "中位", "風力大"]:
        r = g.loc[ter]
        rows2.append({"季節": name.split("（")[0], "風力三分位": ter,
                      "TB4hスプレッド(円/kWh)": round(r["TB4hスプレッド"], 2),
                      "PF粗利(円/kW-日)": round(r["PF粗利"] / 1000, 1),
                      "床時間(h/日)": round(r["床時間"], 2),
                      "安値分断(h/日)": round(r["安値分断時間"], 1)})
t2 = pd.DataFrame(rows2)
print("\n=== 風力三分位×裁定指標 ===")
print(t2.to_string(index=False))

# ============================================================
# 3. duration曲線: PF価値のh依存（2/4/6/8h）
# ============================================================
dur_rows = []
for h in (2, 4, 6, 8):
    m = np.maximum(0.0, (srt[:, -h:].sum(axis=1) * ETA - srt[:, :h].sum(axis=1)) * 1000)
    fy = P.index.year - (P.index.month < 4).astype(int)
    ann = pd.Series(m, index=P.index).groupby(fy).sum() / 1000 / 1  # 円/kW-年（1MW=1000kW, MWh→kWh係数と相殺）
    dur_rows.append({"時間率": f"{h}h", "PF価値(円/kW-年, FY23-25平均)": round(ann.mean(), 0),
                     "4h比": round(ann.mean() / None if h == 0 else 0, 2)})
# 4h比を後計算
base4 = [r for r in dur_rows if r["時間率"] == "4h"][0]["PF価値(円/kW-年, FY23-25平均)"]
for r in dur_rows:
    r["4h比"] = round(r["PF価値(円/kW-年, FY23-25平均)"] / base4, 2)
t3 = pd.DataFrame(dur_rows)
print("\n=== duration曲線（スポットPF価値） ===")
print(t3.to_string(index=False))

# ============================================================
# 4. 風力の日内プロファイルと抑制実績
# ============================================================
prof = {}
for name, months in SEASONS:
    sub = p[p["ts"].dt.month.isin(months)]
    prof[name.split("（")[0]] = sub.groupby("hour")["wind_pre"].mean()
t4 = pd.DataFrame(prof)
flat = (t4.max() - t4.min()) / t4.mean() * 100
print("\n風力の日内振幅（(max-min)/mean %）:", flat.round(1).to_dict())
full = pd.read_csv(os.path.join(PROC, "hokkaido_hourly_panel.csv"), parse_dates=["ts"])
full["fy"] = full["ts"].dt.year - (full["ts"].dt.month < 4).astype(int)
curt = full.groupby("fy").apply(lambda g: g["wind_curt"].fillna(0).sum()
                                / max(1.0, g["wind"].fillna(0).sum() + g["wind_curt"].fillna(0).sum()) * 100)
print("風力抑制率（%・エネルギーベース）:", curt.loc[2022:].round(2).to_dict())

# ============================================================
# Excel追記
# ============================================================
wb = load_workbook(XLSX)
HDR = PatternFill("solid", fgColor="0079C2")


def style_hdr(ws, row, c0, n):
    for c in range(c0, c0 + n):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")


def put_df(ws, df, r0, c0):
    for k, col in enumerate(df.columns):
        ws.cell(row=r0, column=c0 + k, value=col)
    style_hdr(ws, r0, c0, len(df.columns))
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for k, v in enumerate(row):
            ws.cell(row=r0 + i, column=c0 + k, value=v)
    return r0, r0 + len(df)


for sheet in ["風力価格カーブ", "風力×裁定指標", "duration曲線"]:
    if sheet in wb.sheetnames:
        del wb[sheet]

# 風力価格カーブ
ws = wb.create_sheet("風力価格カーブ")
ws["A1"] = "風力大/小の日の時間帯別価格カーブ（3季節、FY2023-25、円/kWh）"
ws["A1"].font = Font(bold=True, size=13, color="176871")
anchors = ["A31", "H31", "O31"]
for si, (name, _) in enumerate(SEASONS):
    cv = curves[name].reset_index().rename(columns={"hour": "時刻"})
    c0 = 1 + si * 3
    ws.cell(row=3, column=c0, value=name).font = Font(bold=True, size=11)
    put_df(ws, cv.round(2), 4, c0)
    ch = LineChart()
    ch.title = name
    ch.height, ch.width = 8.5, 11.5
    data = Reference(ws, min_col=c0 + 1, max_col=c0 + 2, min_row=4, max_row=28)
    cats = Reference(ws, min_col=c0, min_row=5, max_row=28)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    for ln, color in zip(ch.series, ["D95B20", "0079C2"]):
        ln.graphicalProperties.line.solidFill = color
        ln.graphicalProperties.line.width = 22000
    ws.add_chart(ch, anchors[si])
ws["A29"] = f"注: 風力大/小＝各季節内で日次風力発電量（制御前）の上位/下位1/3の日（日数: " + " / ".join(
    f"{n.split('（')[0]} {a}:{b}" for n, (a, b) in counts.items()) + "）"
ws["A29"].font = Font(size=9, color="595959")

# 風力×裁定指標
ws = wb.create_sheet("風力×裁定指標")
ws["A1"] = "風力の日次水準（三分位）× 裁定指標（FY2023-25・日次平均）"
ws["A1"].font = Font(bold=True, size=13, color="176871")
put_df(ws, t2, 3, 1)
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 11
for col in "CDEF":
    ws.column_dimensions[col].width = 19
ch = BarChart()
ch.type = "col"
ch.title = "PF粗利（円/kW-日）: 風力三分位×季節"
ch.height, ch.width = 9, 15
# データを季節×三分位で並べ替えた小テーブルを右に作る
pv = t2.pivot(index="風力三分位", columns="季節", values="PF粗利(円/kW-日)").reindex(["風力小", "中位", "風力大"])
r0, r1 = put_df(ws, pv.reset_index(), 3, 8)
data = Reference(ws, min_col=9, max_col=11, min_row=r0, max_row=r1)
cats = Reference(ws, min_col=8, min_row=r0 + 1, max_row=r1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ws.add_chart(ch, "H9")
ws["A14"] = "注: TB4h=上位4h平均−下位4h平均。PF粗利=完全予見4h・η0.85。安値分断=道内<システム−0.01円の時間"
ws["A14"].font = Font(size=9, color="595959")

# duration曲線
ws = wb.create_sheet("duration曲線")
ws["A1"] = "duration曲線: 完全予見裁定価値のh依存（実績価格、FY2023-25平均）＋容量市場κ(h)"
ws["A1"].font = Font(bold=True, size=13, color="176871")
t3x = t3.copy()
t3x["容量市場κ(h) 2029年度"] = ["—", "83.6%", "93.2%", "98.3%"]
put_df(ws, t3x, 3, 1)
for col in "ABCD":
    ws.column_dimensions[col].width = 24
ch = BarChart()
ch.type = "col"
ch.title = "スポットPF価値（円/kW-年）"
ch.height, ch.width = 8.5, 12
data = Reference(ws, min_col=2, min_row=3, max_row=7)
cats = Reference(ws, min_col=1, min_row=4, max_row=7)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ch.legend = None
ws.add_chart(ch, "F3")
ws["A9"] = "注: 出力1kWあたり。hを増やすと総価値は増えるが限界価値は逓減（浅い谷・鋭いピークの形状による）。κはOCCTO調整係数（北海道・蓄電池）"
ws["A9"].font = Font(size=9, color="595959")

toc = wb["目次"]
for sheet, desc in [("風力価格カーブ", "追加: 風力大小×3季節の価格カーブ"),
                    ("風力×裁定指標", "追加: 風力三分位×TB4h/PF/床/安値分断"),
                    ("duration曲線", "追加: PF価値のh依存＋κ(h)")]:
    row = toc.max_row + 1
    toc.cell(row=row, column=1, value=sheet)
    toc.cell(row=row, column=2, value=desc)
wb.save(XLSX)
print("saved xlsx sheets")
