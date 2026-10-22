#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""K_windスイープ: 風力導入量を0.5×〜3×に振ったときの蓄電池スポット価値（価格過程v1）

13_price_process_v1 のモデル一式（cf過程・g+μ・分断3レジーム混合・共通乱数）を実行し、
K_wind乗数ごとに 床時間・TB4h・PF価値 を再計算する。泊3号再稼働の有無も併記。
これが「風力増加×蓄電池価値」の中心図: 現行水準近傍では微増、床飽和後に減少へ転じる転換点を探す。

出力: figures/hokkaido/kwind_sweep.png ＋ 図表データ.xlsx「風力スイープ」シート
"""
import os

import numpy as np
import pandas as pd
from matplotlib import font_manager
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

for f in ["/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
          "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"]:
    try:
        font_manager.fontManager.addfont(f)
    except Exception:
        pass
plt.rcParams["font.family"] = ["Noto Sans CJK JP", "Hiragino Sans", "Yu Gothic", "Meiryo"]

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.join(HERE, "..")
FIGDIR = os.path.join(ROOT, "figures", "hokkaido")
XLSX = os.path.join(ROOT, "slides", "進捗報告_20260817_図表データ.xlsx")

# 13を実行してモデル部品（est, simulate_B, metrics 等）を取り込む
ns = {"__file__": os.path.join(HERE, "13_price_process_v1.py"), "__name__": "pp_v1"}
exec(open(os.path.join(HERE, "13_price_process_v1.py"), encoding="utf-8").read(), ns)
est = ns["est"]
simulate_B = ns["simulate_B"]
metrics = ns["metrics"]
K_WIND_NOW = ns["K_WIND_NOW"]  # 万kW

MULTS = [0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0, 2.5, 3.0]
TOMARI = 912_000 * 0.9 / 1000  # MW

rows = []
for tomari in (0, TOMARI):
    for m in MULTS:
        sc = est.copy()
        dwind_kw = (m - 1.0) * K_WIND_NOW * 1e4
        sc["net"] = sc["net"] - dwind_kw * sc["cf_wind"] / 1000 - tomari
        sc["p_B"] = simulate_B(sc)
        mt = metrics(sc, "p_B")
        rows.append({
            "泊": "あり" if tomari else "なし",
            "K_wind倍率": m,
            "K_wind(万kW)": round(K_WIND_NOW * m, 0),
            "床時間(h/年)": round(mt["床(≤0.01円)時間"].mean(), 0),
            "TB4h中央値(円/kWh)": round(mt["TB4h中央値"].mean(), 2),
            "PF価値(円/kW-年)": round(mt["PF価値(円/kW-年)"].mean(), 0),
        })
        print(rows[-1])

df = pd.DataFrame(rows)
df.to_csv(os.path.join(ROOT, "data", "processed", "kwind_sweep.csv"), index=False)

# ---------- PNG（上: 完全予見粗利、下: 下限張り付き時間。横軸を共有して縦に並べる） ----------
# 論文の Word 版では幅15〜15.5cm に縮小して載せるので、幅6.1インチ（15.5cm）で作り、
# 図中の文字を9〜10pt にする（縮小後も8pt 以上）
FIG_W, FS_TICK, FS_LABEL, FS_TITLE = 6.1, 9, 9.5, 10
C0, C1, GRAY = "#0079C2", "#D95B20", "#595959"
fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(FIG_W, 5.0), sharex=True, layout="constrained")
for tomari, color, lab, ls, mk in [("なし", C0, "泊なし（現状）", "-", "o"), ("あり", C1, "泊3号再稼働", "--", "s")]:
    d = df[df["泊"] == tomari]
    ax1.plot(d["K_wind(万kW)"], d["PF価値(円/kW-年)"], color=color, lw=1.8, ls=ls, marker=mk, ms=4, label=lab)
    ax2.plot(d["K_wind(万kW)"], d["床時間(h/年)"], color=color, lw=1.8, ls=ls, marker=mk, ms=4, label=lab)
for ax, ttl, yl in [(ax1, "完全予見粗利（分断混合仕様）", "円/kW-年"),
                    (ax2, "下限張り付き時間（価格＝0.01円/kWh）", "時間/年")]:
    ax.set_title(ttl, fontsize=FS_TITLE, pad=4)
    ax.set_ylabel(yl, fontsize=FS_LABEL, color=GRAY)
    ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda v, _: f"{v:,.0f}"))
    ax.grid(axis="y", color="#DDDDDD", lw=0.6)
    for sp in ["top", "right"]:
        ax.spines[sp].set_visible(False)
    ax.tick_params(colors=GRAY, labelsize=FS_TICK)
    ax.axvline(K_WIND_NOW, color="#999999", lw=1, ls="--")
ax2.set_xlabel("風力導入量（万kW）", fontsize=FS_LABEL, color=GRAY)
# 「現状」は破線の上端に置き、凡例は図の上に出して系列と重ならないようにする
ax1.annotate("現状", (K_WIND_NOW, ax1.get_ylim()[1]), xytext=(4, -2), textcoords="offset points",
             va="top", fontsize=FS_TICK, color=GRAY)
fig.legend(*ax1.get_legend_handles_labels(), fontsize=FS_TICK, frameon=False, ncol=2, loc="outside upper center")
# 図題・注・出所は論文本文（図の下）に置くので、画像には入れない
fig.savefig(os.path.join(FIGDIR, "kwind_sweep.png"), dpi=300)
print("saved kwind_sweep.png")

# ---------- Excel ----------
wb = load_workbook(XLSX)
if "風力スイープ" in wb.sheetnames:
    del wb["風力スイープ"]
ws = wb.create_sheet("風力スイープ")
ws["A1"] = "K_windスイープ: 風力導入量×蓄電池スポット価値（価格過程v1、FY2023-25平均）"
ws["A1"].font = Font(bold=True, size=13, color="176871")
HDR = PatternFill("solid", fgColor="0079C2")
for k, col in enumerate(df.columns, start=1):
    cell = ws.cell(row=3, column=k, value=col)
    cell.fill = HDR
    cell.font = Font(bold=True, color="FFFFFF")
    cell.alignment = Alignment(horizontal="center")
for i, (_, row) in enumerate(df.iterrows(), start=4):
    for k, v in enumerate(row, start=1):
        ws.cell(row=i, column=k, value=v)
for col, w in zip("ABCDEF", [8, 12, 14, 14, 18, 18]):
    ws.column_dimensions[col].width = w

npts = len(MULTS)
ch = LineChart()
ch.title = "PF価値（円/kW-年） vs K_wind（万kW）"
ch.height, ch.width = 10, 16
for bi, (label, r0) in enumerate([("泊なし", 4), ("泊あり", 4 + npts)]):
    data = Reference(ws, min_col=6, min_row=r0, max_row=r0 + npts - 1)
    ch.add_data(data, titles_from_data=False)
    ch.series[bi].tx = None
cats = Reference(ws, min_col=3, min_row=4, max_row=3 + npts)
ch.set_categories(cats)
for ln, color in zip(ch.series, ["0079C2", "D95B20"]):
    ln.graphicalProperties.line.solidFill = color
    ln.graphicalProperties.line.width = 22000
    ln.smooth = False
ch.x_axis.delete = False
ch.y_axis.delete = False
ws.add_chart(ch, "H3")
ws["A25"] = "注: 部分均衡（蓄電池フリート応答・p_system・供給側の内生反応は未反映）。青=泊なし/橙=泊あり。K_wind現状121万kW"
ws["A25"].font = Font(size=9, color="595959")

toc = wb["目次"]
row = toc.max_row + 1
toc.cell(row=row, column=1, value="風力スイープ")
toc.cell(row=row, column=2, value="追加: K_wind 0.5-3×の蓄電池価値カーブ（泊あり/なし）")
wb.save(XLSX)
print("saved xlsx 風力スイープ")
