#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""帯域分解の比較: 北海道風力 vs 九州太陽光（いずれも制御前出力）

  1. 静的比較（FY2022-25プール）: 北海道風力・九州太陽光＋参考（九州風力・北海道太陽光）
  2. 九州太陽光の年度別時系列（FY2016-25）— 北海道風力版（19）との対
出力: figures/hokkaido/band_comparison.png ＋ 図表データ.xlsx「帯域分解比較」シート
"""
import os

import pandas as pd
from matplotlib import font_manager
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

for f in ["/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
          "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc"]:
    try:
        font_manager.fontManager.addfont(f)
    except Exception:
        pass
plt.rcParams["font.family"] = "Noto Sans CJK JP"

ROOT = os.path.join(os.path.dirname(__file__), "..")
PROC = os.path.join(ROOT, "data", "processed")
FIGDIR = os.path.join(ROOT, "figures", "hokkaido")
XLSX = os.path.join(ROOT, "slides", "進捗報告_20260817_図表データ.xlsx")
GRAY = "#595959"
BANDS = ["<6時間", "6-24時間", "1-7日", ">7日"]


def band_shares(x):
    x = x.interpolate()
    ma6 = x.rolling(6, center=True, min_periods=3).mean()
    ma24 = x.rolling(24, center=True, min_periods=12).mean()
    ma168 = x.rolling(168, center=True, min_periods=84).mean()
    comps = {"<6時間": x - ma6, "6-24時間": ma6 - ma24,
             "1-7日": ma24 - ma168, ">7日": ma168 - x.mean()}
    tot = x.var()
    return {k: float(v.var() / tot * 100) for k, v in comps.items()}


hp = pd.read_csv(os.path.join(PROC, "hokkaido_hourly_panel.csv"), parse_dates=["ts"])
kp = pd.read_csv(os.path.join(PROC, "kyushu_hourly_panel.csv"), parse_dates=["ts"])
for df in (hp, kp):
    df["fy"] = df["ts"].dt.year - (df["ts"].dt.month < 4).astype(int)
    df["wind_pre"] = df["wind"].fillna(0) + df["wind_curt"].fillna(0)
    df["solar_pre"] = df["solar"].fillna(0) + df["solar_curt"].fillna(0)

# ---------- 1. 静的比較（FY2022-25） ----------
sel_h = hp[(hp["ts"] >= "2022-04-01") & (hp["ts"] < "2026-04-01")].set_index("ts")
sel_k = kp[(kp["ts"] >= "2022-04-01") & (kp["ts"] < "2026-04-01")].set_index("ts")
static = pd.DataFrame({
    "北海道 風力": band_shares(sel_h["wind_pre"]),
    "九州 太陽光": band_shares(sel_k["solar_pre"]),
    "九州 風力": band_shares(sel_k["wind_pre"]),
    "北海道 太陽光": band_shares(sel_h["solar_pre"]),
}).loc[BANDS]
print("=== 静的比較（FY2022-25、分散シェア%） ===")
print(static.round(1).to_string())

# ---------- 2. 九州太陽光の年度別時系列 ----------
rows = []
for fy in range(2016, 2026):
    g = kp[kp["fy"] == fy].set_index("ts")["solar_pre"]
    if g.notna().sum() < 5000:
        continue
    sh = band_shares(g)
    sh["fy"] = fy
    rows.append(sh)
ks = pd.DataFrame(rows).set_index("fy")[BANDS]
print("\n=== 九州太陽光の年度別時系列 ===")
print(ks.round(1).to_string())

# ---------- PNG ----------
C_HW, C_KS = "#176871", "#D95B20"
C_BANDS = {"<6時間": "#5DB6E7", "6-24時間": "#0079C2", "1-7日": "#176871", ">7日": "#D95B20"}
fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(12.8, 4.8), gridspec_kw={"width_ratios": [1, 1.25]})

xpos = np.arange(len(BANDS))
w = 0.36
ax1.bar(xpos - w / 2, static["北海道 風力"], w, color=C_HW, label="北海道 風力")
ax1.bar(xpos + w / 2, static["九州 太陽光"], w, color=C_KS, label="九州 太陽光")
for i, b in enumerate(BANDS):
    ax1.text(i - w / 2, static.loc[b, "北海道 風力"] + 1, f"{static.loc[b, '北海道 風力']:.0f}", ha="center", fontsize=9.5, color=C_HW, fontweight="bold")
    ax1.text(i + w / 2, static.loc[b, "九州 太陽光"] + 1, f"{static.loc[b, '九州 太陽光']:.0f}", ha="center", fontsize=9.5, color=C_KS, fontweight="bold")
ax1.set_xticks(xpos)
ax1.set_xticklabels(BANDS, fontsize=10)
ax1.set_ylabel("分散シェア（%）", fontsize=10, color=GRAY)
ax1.set_title("帯域分解の対比（FY2022-25）", fontsize=12.5)
ax1.legend(fontsize=10, frameon=False, loc="upper left")

for col, color in C_BANDS.items():
    ax2.plot(ks.index, ks[col], color=color, lw=2.2, marker="o", ms=4.5, label=col)
ax2.set_title("九州 太陽光: 帯域分解の推移（年度別）", fontsize=12.5)
ax2.set_xticks(ks.index)
ax2.set_xticklabels([f"'{str(y)[2:]}" for y in ks.index], fontsize=9.5)
ax2.set_ylabel("分散シェア（%）", fontsize=10, color=GRAY)
ax2.set_ylim(0, None)
ax2.legend(fontsize=9.5, frameon=False, ncol=2, loc="center left", bbox_to_anchor=(0.05, 0.5))

for ax in (ax1, ax2):
    ax.grid(axis="y", color="#DDDDDD", lw=0.6)
    for sp in ["top", "right"]:
        ax.spines[sp].set_visible(False)
    ax.tick_params(colors=GRAY, labelsize=9.5)
fig.suptitle("北海道風力は長周期（1日超が約3/4）、九州太陽光は日内（6-24時間が約2/3）— 構造は年度を通じて安定",
             fontsize=13, fontweight="bold", y=1.02)
fig.text(0.995, -0.05,
         "いずれも制御前出力（出力+抑制量）・フリート集約。移動平均カスケードによる非直交分解（交差項ありシェア合計≠100%）",
         ha="right", fontsize=8, color=GRAY)
fig.tight_layout()
fig.savefig(os.path.join(FIGDIR, "band_comparison.png"), dpi=160, bbox_inches="tight")
print("saved band_comparison.png")

# ---------- Excel ----------
wb = load_workbook(XLSX)
if "帯域分解比較" in wb.sheetnames:
    del wb["帯域分解比較"]
ws = wb.create_sheet("帯域分解比較")
ws["A1"] = "帯域分解の比較: 北海道風力 vs 九州太陽光（制御前出力、移動平均カスケード）"
ws["A1"].font = Font(bold=True, size=13, color="176871")
HDR = PatternFill("solid", fgColor="0079C2")


def put(ws, df, r0, title, index_name):
    ws.cell(row=r0, column=1, value=title).font = Font(bold=True, size=11)
    r0 += 1
    heads = [index_name] + list(df.columns)
    for k, h in enumerate(heads, start=1):
        cell = ws.cell(row=r0, column=k, value=str(h))
        cell.fill = HDR
        cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center")
    for i, (idx, row) in enumerate(df.iterrows(), start=1):
        ws.cell(row=r0 + i, column=1, value=str(idx))
        for k, v in enumerate(row, start=2):
            ws.cell(row=r0 + i, column=k, value=round(float(v), 1))
    return r0, r0 + len(df)


h0, h1 = put(ws, static, 3, "① 静的比較（FY2022-25プール、分散シェア%）", "帯域")
ch = BarChart()
ch.type = "col"
ch.title = "北海道風力 vs 九州太陽光（分散シェア%）"
ch.height, ch.width = 8.5, 13
data = Reference(ws, min_col=2, max_col=3, min_row=h0, max_row=h1)
cats = Reference(ws, min_col=1, min_row=h0 + 1, max_row=h1)
ch.add_data(data, titles_from_data=True)
ch.set_categories(cats)
ws.add_chart(ch, "G3")

k0, k1 = put(ws, ks, h1 + 3, "② 九州太陽光の年度別時系列（分散シェア%）", "年度")
ch2 = LineChart()
ch2.title = "九州太陽光: 帯域別分散シェアの推移"
ch2.height, ch2.width = 8.5, 13
data = Reference(ws, min_col=2, max_col=5, min_row=k0, max_row=k1)
cats = Reference(ws, min_col=1, min_row=k0 + 1, max_row=k1)
ch2.add_data(data, titles_from_data=True)
ch2.set_categories(cats)
for ln, color in zip(ch2.series, ["5DB6E7", "0079C2", "176871", "D95B20"]):
    ln.graphicalProperties.line.solidFill = color
    ln.graphicalProperties.line.width = 22000
    ln.smooth = False
ch2.x_axis.delete = False
ch2.y_axis.delete = False
ws.add_chart(ch2, "G20")
for col, w_ in zip("ABCDE", [12, 13, 13, 13, 13]):
    ws.column_dimensions[col].width = w_
ws.cell(row=k1 + 2, column=1, value="注: いずれも制御前出力（出力+抑制量）。非直交分解のためシェア合計≠100%。①の参考列: 九州風力・北海道太陽光").font = Font(size=9, color="595959")

toc = wb["目次"]
row = toc.max_row + 1
toc.cell(row=row, column=1, value="帯域分解比較")
toc.cell(row=row, column=2, value="追加: 北海道風力 vs 九州太陽光の帯域分解（静的＋九州太陽光の時系列）")
wb.save(XLSX)
print("saved xlsx 帯域分解比較")
