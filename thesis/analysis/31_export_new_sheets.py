#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""2026-09-29 に追加した解析結果を図表データ.xlsx にシートとして書き出す（冪等: 既存シートは置換、目次は重複しない）

対象: 価格水準θ（26）、π(K)水準感応度（27）、均衡面・フロンティア・K*格子（28）、空間分散（29）、BTM回避率（30）、燃料危機期除外回帰（24）
"""
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
PROC = os.path.join(ROOT, "data", "processed")
XLSX = os.path.join(ROOT, "slides", "進捗報告_20260817_図表データ.xlsx")
HDR = PatternFill("solid", fgColor="0079C2")
COLORS = ["0079C2", "D95B20", "176871", "7B3F9E", "9DC3E6"]


def put_df(ws, df, r0, c0=1, title=None):
    if title:
        ws.cell(row=r0, column=c0, value=title).font = Font(bold=True, size=11)
        r0 += 1
    for k, col in enumerate(df.columns):
        cell = ws.cell(row=r0, column=c0 + k, value=str(col))
        cell.fill = HDR; cell.font = Font(bold=True, color="FFFFFF", size=9); cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (_, row) in enumerate(df.iterrows(), start=1):
        for k, v in enumerate(row):
            ws.cell(row=r0 + i, column=c0 + k, value=(None if pd.isna(v) else (float(v) if hasattr(v, "item") else v)))
    return r0, r0 + len(df)


def line_chart(ws, title, hdr_row, last_row, cat_col, data_cols, anchor, ytitle=""):
    ch = LineChart(); ch.title = title; ch.height, ch.width = 8.5, 13
    for i, c in enumerate(data_cols):
        ch.add_data(Reference(ws, min_col=c, min_row=hdr_row, max_row=last_row), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=cat_col, min_row=hdr_row + 1, max_row=last_row))
    for ln, color in zip(ch.series, COLORS * 3):
        ln.graphicalProperties.line.solidFill = color; ln.graphicalProperties.line.width = 22000; ln.smooth = False
    ch.x_axis.delete = False; ch.y_axis.delete = False; ch.y_axis.title = ytitle
    ws.add_chart(ch, anchor)


wb = load_workbook(XLSX)
sheets = {}

# 1. 価格水準θ
th = pd.read_csv(os.path.join(PROC, "price_level_theta.csv"))
ws = wb.create_sheet("価格水準θ") if "価格水準θ" not in wb.sheetnames else (wb.remove(wb["価格水準θ"]) or wb.create_sheet("価格水準θ"))
ws["A1"] = "供給曲線v2の価格水準係数 θ（FY2023-25 の時間加重平均=1）: 年度別と月次"; ws["A1"].font = Font(bold=True, size=13, color="176871")
fy = th[th["ym"].str.startswith("FY")][["ym", "theta_m", "mean_p_hokkaido", "mean_p_system"]].rename(columns={"ym": "年度", "theta_m": "θ_FY", "mean_p_hokkaido": "平均価格(北海道)", "mean_p_system": "平均システムプライス"})
r0, r1 = put_df(ws, fy, 3, 1, "年度別 θ")
mo = th[~th["ym"].str.startswith("FY")][["ym", "theta_m", "mean_p_hokkaido", "mean_p_system"]].rename(columns={"ym": "年月", "theta_m": "θ_m", "mean_p_hokkaido": "平均価格(北海道)", "mean_p_system": "平均システムプライス"})
r0m, r1m = put_df(ws, mo, 3, 7, "月次 θ_m（価格水準指数）")
line_chart(ws, "月次 θ_m", r0m, r1m, 7, [8], "M3", "θ")
sheets["価格水準θ"] = "追加(9/29): 供給曲線v2の年度別・月次の価格水準係数θ（表8.2）"

# 2. π(K) 水準感応度
pk = pd.read_csv(os.path.join(PROC, "pi_k_curve_v2.csv"))
name = "π(K)水準感応度"
if name in wb.sheetnames: wb.remove(wb[name])
ws = wb.create_sheet(name)
ws["A1"] = "π(K) の価格水準感応度（供給曲線v2、円/kW-年、括弧なし=モデル値、_裾補正=×1.25）"; ws["A1"].font = Font(bold=True, size=13, color="176871")
pv = pk.pivot(index="K(万kW)", columns="シナリオ", values="π_実現可能(円/kW-年)").reset_index()
r0, r1 = put_df(ws, pv, 3, 1, "π_実現可能(K)")
pv2 = pk.pivot(index="K(万kW)", columns="シナリオ", values="π_実現可能_裾補正(円/kW-年)").reset_index()
put_df(ws, pv2, r1 + 3, 1, "π_実現可能_裾補正(K)")
line_chart(ws, "π_実現可能(K) シナリオ別", r0, r1, 1, list(range(2, 2 + len(pv.columns) - 1)), "H3", "円/kW-年")
be = pd.read_csv(os.path.join(PROC, "pi_k_breakeven_theta.csv"))
put_df(ws, be, r1 + 14, 1, "損益分岐の θ*（K=0）")
sheets[name] = "追加(9/29): θ=1／FY2026上期／FY2022／泊 の π(K) と損益分岐θ*（表8.5・図8.3）"

# 3. 均衡面
es = pd.read_csv(os.path.join(PROC, "equilibrium_surface.csv"))
name = "均衡面"
if name in wb.sheetnames: wb.remove(wb[name])
ws = wb.create_sheet(name)
ws["A1"] = "均衡面 π(K; K_wind, 泊)（供給曲線v2 θ=1、円/kW-年）"; ws["A1"].font = Font(bold=True, size=13, color="176871")
r = 3
for tomari in ("なし", "あり"):
    pv = es[es["泊"] == tomari].pivot(index="K(万kW)", columns="K_wind倍率", values="π_実現可能(円/kW-年)").reset_index()
    pv.columns = [pv.columns[0]] + [f"風力×{c}" for c in pv.columns[1:]]
    r0, r1 = put_df(ws, pv, r, 1, f"泊{tomari}")
    line_chart(ws, f"π(K; K_wind) 泊{tomari}", r0, r1, 1, list(range(2, 2 + len(pv.columns) - 1)), "I" + str(r), "円/kW-年")
    r = r1 + 3
sheets[name] = "追加(9/29): 風力0.5〜3×の均衡面（表8.6・図8.4）"

# 4. フロンティア
fr = pd.read_csv(os.path.join(PROC, "breakeven_frontier.csv"))
ks = pd.read_csv(os.path.join(PROC, "kstar_grid.csv"))
cp = pd.read_csv(os.path.join(PROC, "capacity_price_of_k.csv"))
name = "フロンティア"
if name in wb.sheetnames: wb.remove(wb[name])
ws = wb.create_sheet(name)
ws["A1"] = "break-even フロンティア: θ*(資本費, 実効容量収入)・K* 格子・P_cap(K)"; ws["A1"].font = Font(bold=True, size=13, color="176871")
r = 3
for w in (0.05, 0.06, 0.08):
    pv = fr[fr["WACC"] == w].pivot(index="資本費(万円/kWh)", columns="実効容量収入", values="θ*").reset_index()
    r0, r1 = put_df(ws, pv, r, 1, f"θ*（WACC {int(w*100)}%）"); r = r1 + 2
r0, r1 = put_df(ws, ks, r, 1, "K* 格子（万kW、WACC6%）"); r = r1 + 2
put_df(ws, cp, r, 1, "容量市場価格の K 依存 P_cap(K)")
sheets[name] = "追加(9/29): θ*格子・K*格子・P_cap(K)（表9.2・9.3・8.7、図9.1）"

# 5. 空間分散
sd = pd.read_csv(os.path.join(PROC, "spatial_dispersion.csv"))
name = "空間分散"
if name in wb.sheetnames: wb.remove(wb[name])
ws = wb.create_sheet(name)
ws["A1"] = "空間分散シミュレーション: 日内帯域の振幅倍率 λ と価格指標・風力→TB4h係数"; ws["A1"].font = Font(bold=True, size=13, color="176871")
r0, r1 = put_df(ws, sd, 3, 1)
line_chart(ws, "PF価値 vs λ", r0, r1, 1, [sd.columns.get_loc("PF価値(円/kW-年)") + 1], "A12", "円/kW-年")
line_chart(ws, "風力→TB4h 係数 vs λ", r0, r1, 1, [sd.columns.get_loc(c) + 1 for c in sd.columns if c.startswith("β_wind")], "I12", "円/kWh per GWh/日")
sheets[name] = "追加(9/29): 集中立地の反実仮想（表6.4・図6.5）"

# 6. BTM 回避率
bt = pd.read_csv(os.path.join(PROC, "btm_avoidable_share.csv"))
name = "BTM回避率"
if name in wb.sheetnames: wb.remove(wb[name])
ws = wb.create_sheet(name)
ws["A1"] = "併設蓄電池の容量制約下での抑制回避可能率（北海道風力、2025/4〜2026/6）"; ws["A1"].font = Font(bold=True, size=13, color="176871")
put_df(ws, bt, 3, 1)
ws["A10"] = "注: 集中=容量シェア20%のサイトに抑制の77%（Maji et al. 2025 の ERCOT 分布を援用）。フリート集約 SoC シミュレーション（analysis/30）"; ws["A10"].font = Font(size=9, color="595959")
sheets[name] = "追加(9/29): 回避可能率 一様／集中（表9.1）"

# 7. 燃料危機期除外
m = pd.read_csv(os.path.join(PROC, "vol_regression_results.csv")); e = pd.read_csv(os.path.join(PROC, "vol_regression_results_ex21-22.csv"))
j = m.merge(e, on=["エリア", "被説明変数", "係数"], suffixes=("_主分析", "_除外"))[["エリア", "被説明変数", "係数", "推定値_主分析", "p値_主分析", "推定値_除外", "p値_除外"]]
name = "燃料危機期除外"
if name in wb.sheetnames: wb.remove(wb[name])
ws = wb.create_sheet(name)
ws["A1"] = "日次回帰: 主分析（FY2016-25）と燃料危機期（FY2021-22）除外の係数比較"; ws["A1"].font = Font(bold=True, size=13, color="176871")
put_df(ws, j.round(4), 3, 1)
sheets[name] = "追加(9/29): 6.5.4 の頑健性（表6.8）"

toc = wb["目次"]
existing = {toc.cell(row=r, column=1).value for r in range(1, toc.max_row + 1)}
for sheet, desc in sheets.items():
    if sheet in existing:
        continue
    row = toc.max_row + 1
    toc.cell(row=row, column=1, value=sheet); toc.cell(row=row, column=2, value=desc)
wb.save(XLSX)
print("saved sheets:", list(sheets))
