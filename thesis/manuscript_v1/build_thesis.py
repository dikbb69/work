#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""修士論文（正式版の体裁）の Word ファイルを作る

入力: manuscript_v1/ の Markdown（ファイル名の数字の順に結合。90_ 以降は別冊として除く）
出力: manuscript_v1/thesis_<tag>.docx（本体）、manuscript_v1/abstract_<tag>.docx（別冊要旨）、結合版 Markdown

体裁（scratchpad/format_spec.md の仕様に従う）:
  - A4 縦、余白 上30・下25・左30・右25mm、本文 游明朝 10.5pt・行送り18pt（約40字×35行）、見出し 游ゴシック
  - 中表紙（頁番号なし）→ 前付け（要旨・Abstract・目次・図目次・表目次・略語・記号。頁番号 i, ii…）→ 本文以降（1, 2…）
  - 目次（3階層）・図目次・表目次は Word のフィールド（開いたときに更新を求める設定）
  - 表題は表の上（pandoc の表題行）、図題は図の下、注）・出所）は図表の直後の段落（9pt）
  - 表は横罫だけの三線表（上端・見出しの下・下端）
  - 別行の式の直後の段落「（式番号 5.1）」を、式を中央・番号 (5.1) を右寄せにした1行に変換する
  - 各ファイルの先頭で改ページ
Markdown 側の約束: 目次等の挿入位置は <!-- TOC --> <!-- LOF --> <!-- LOT -->、前付けの見出しは `::: {custom-style="Front Heading"}` で囲む。
  参考文献の項目は `::: {custom-style="Bibliography"}` で囲む（ぶら下げインデントの書式が付く）。
使い方: python build_thesis.py [--tag draft]
"""
import argparse
import copy
import glob
import os
import re
import tempfile

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
PAGEBREAK = '\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\n'
SECTION_MARK = "〔SECTION-BREAK:{name}〕"


def field_xml(instr):
    return ('```{=openxml}\n<w:p><w:r><w:fldChar w:fldCharType="begin" w:dirty="true"/></w:r>'
            f'<w:r><w:instrText xml:space="preserve"> {instr} </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r><w:r><w:t>（Word で開いたときに更新されます）</w:t></w:r>'
            '<w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>\n```')


def fix_math(text):
    """数式中の日本語を \\text{} で包み、命令と日本語が続く箇所に空白を入れる（既存の \\text{} は保つ）"""
    def conv(seg):
        seg = re.sub(r"(\\[A-Za-z]+)(?=[^\x00-\x7F])", r"\1 ", seg)
        return re.sub(r"([^\x00-\x7F\s][^\x00-\x7F]*)", lambda k: "\\text{" + k.group(1) + "}", seg)

    def wrap(m):
        parts = re.split(r"(\\(?:text|mathrm|textrm|mbox|operatorname)\{[^{}]*\})", m.group(2))
        return m.group(1) + "".join(p if i % 2 else conv(p) for i, p in enumerate(parts)) + m.group(3)
    text = re.sub(r"(\$\$)(.+?)(\$\$)", wrap, text, flags=re.S)
    text = re.sub(r"(?<![\\$])(\$)([^$\n]+?)(\$)(?!\$)", wrap, text)
    return text


def collect(tag):
    files = sorted(f for f in glob.glob(os.path.join(HERE, "[0-9][0-9]_*.md")))
    main = [f for f in files if int(os.path.basename(f)[:2]) < 90]
    sep = [f for f in files if int(os.path.basename(f)[:2]) >= 90]
    return main, sep


def assemble(files):
    parts = []
    for f in files:
        t = open(f, encoding="utf-8").read().strip()
        name = os.path.basename(f)
        parts.append(t)
        # 区切り: 中表紙の後・前付けの後に節区切り、それ以外は改ページ
        if name.startswith("00_"):
            parts.append(SECTION_MARK.format(name="cover"))
        elif name.startswith("03_"):
            parts.append(SECTION_MARK.format(name="front"))
        else:
            parts.append(PAGEBREAK)
    body = "\n\n".join(parts)
    body = body.replace("<!-- TOC -->", field_xml('TOC \\o "1-3" \\h \\z \\u'))
    body = body.replace("<!-- LOF -->", field_xml('TOC \\h \\z \\t "Image Caption,1"'))
    body = body.replace("<!-- LOT -->", field_xml('TOC \\h \\z \\t "Table Caption,1"'))
    return fix_math(body)


def set_fonts(style, font, size=None, bold=None):
    style.font.name = font
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(k), font)


def line_spacing(style, pts, before=0, after=0):
    pf = style.paragraph_format
    pf.line_spacing = Pt(pts)
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)


def make_reference(path):
    pypandoc.convert_text("x\n\n::: {custom-style=\"Front Heading\"}\nx\n:::\n\n::: {custom-style=\"Cover Title\"}\nx\n:::\n\n::: {custom-style=\"Cover Line\"}\nx\n:::\n\n::: {custom-style=\"Table Note\"}\nx\n:::\n\n::: {custom-style=\"Bibliography\"}\nx\n:::\n",
                          "docx", format="markdown", outputfile=path)
    doc = Document(path)
    st = doc.styles
    MIN, GOT = "Yu Mincho", "Yu Gothic"
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in [s.name for s in st]:
            set_fonts(st[name], MIN, 10.5)
            line_spacing(st[name], 18)
            st[name].paragraph_format.first_line_indent = Pt(10.5) if name in ("Body Text", "First Paragraph") else None
            st[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if name != "Compact" else None
    for name, size, before, after in (("Heading 1", 16, 0, 18), ("Heading 2", 13, 12, 6), ("Heading 3", 11.5, 8, 4)):
        s = st[name]
        set_fonts(s, GOT, size, True)
        line_spacing(s, size * 1.5, before, after)
        s.font.color.rgb = None
        s.paragraph_format.first_line_indent = None
    for name in ("Table Caption", "Image Caption", "Caption"):
        if name in [s.name for s in st]:
            set_fonts(st[name], GOT, 9.5)
            st[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            line_spacing(st[name], 14, 4, 4)
            st[name].font.italic = False
    for name in ("Footnote Text",):
        if name in [s.name for s in st]:
            set_fonts(st[name], MIN, 9)
            line_spacing(st[name], 13, 0, 0)
    s = st["Front Heading"]
    set_fonts(s, GOT, 14, True)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line_spacing(s, 22, 0, 12)
    s = st["Cover Title"]
    set_fonts(s, GOT, 18, True)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line_spacing(s, 30, 12, 12)
    s = st["Cover Line"]
    set_fonts(s, MIN, 12)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    line_spacing(s, 20, 6, 6)
    s = st["Table Note"]
    set_fonts(s, MIN, 9)
    line_spacing(s, 13, 0, 6)
    # 図の段落: Normal の固定行送り（18pt）を継承すると画像が1行の高さで切れるので、行送りを「1行」にして中央に置く
    for name in ("Figure", "Captioned Figure"):
        if name in [x.name for x in st]:
            pf = st[name].paragraph_format
            pf.line_spacing = 1.0
            pf.space_before, pf.space_after = Pt(6), Pt(0)
            pf.first_line_indent = None
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # 東アジアの言語を日本語にする（既定は zh-CN）
    lang = doc.styles.element.find(qn("w:docDefaults") + "/" + qn("w:rPrDefault") + "/" + qn("w:rPr") + "/" + qn("w:lang"))
    if lang is not None:
        lang.set(qn("w:eastAsia"), "ja-JP")
    # 参考文献（40_references.md の Bibliography の div）: 2字のぶら下げ、左揃え（長い URL で字間が開かないように）、項目の間を少し空ける
    s = st["Bibliography"]
    set_fonts(s, MIN, 10.5)
    line_spacing(s, 18, 0, 4)
    s.paragraph_format.left_indent = Pt(21)
    s.paragraph_format.first_line_indent = Pt(-21)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    # 用紙と余白
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(30), Mm(25)
    sec.left_margin, sec.right_margin = Mm(30), Mm(25)
    doc.save(path)


def three_line_tables(doc):
    for tbl in doc.tables:
        tblPr = tbl._tbl.tblPr
        for old in tblPr.findall(qn("w:tblBorders")):
            tblPr.remove(old)
        b = OxmlElement("w:tblBorders")
        for edge, val in (("top", "single"), ("bottom", "single"), ("left", "nil"), ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), val)
            if val == "single":
                e.set(qn("w:sz"), "8")
                e.set(qn("w:color"), "000000")
            b.append(e)
        # スキーマの順序（tblInd の後、shd・tblLayout・tblCellMar・tblLook・tblCaption の前）に入れる
        later = [c for c in tblPr if c.tag in {qn(f"w:{n}") for n in ("shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription")}]
        if later:
            later[0].addprevious(b)
        else:
            tblPr.append(b)
        if tbl.rows:
            for cell in tbl.rows[0].cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tb = OxmlElement("w:tcBorders")
                bot = OxmlElement("w:bottom")
                bot.set(qn("w:val"), "single")
                bot.set(qn("w:sz"), "6")
                bot.set(qn("w:color"), "000000")
                tb.append(bot)
                tcPr.append(tb)
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.first_line_indent = None
                    for r in p.runs:
                        r.font.size = Pt(9)


def number_equations(doc):
    """「（式番号 5.1）」の段落を直前の別行の式と1行にまとめる"""
    body = doc.element.body
    paras = list(body.iter(qn("w:p")))
    usable = Mm(210 - 30 - 25)
    for i, p in enumerate(paras):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        m = re.fullmatch(r"\s*（式番号\s*([0-9A-Z]+\.[0-9]+)）\s*", text)
        if not m or i == 0:
            continue
        prev = paras[i - 1]
        omp = prev.find(".//" + qn("m:oMathPara"))
        if omp is None:
            continue
        om = omp.find(qn("m:oMath"))
        parent = omp.getparent()
        idx = list(parent).index(omp)
        parent.remove(omp)
        tab1 = OxmlElement("w:r"); tab1.append(OxmlElement("w:tab"))
        tab2 = OxmlElement("w:r"); tab2.append(OxmlElement("w:tab"))
        num = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = f"({m.group(1)})"; num.append(t)
        for k, el in enumerate((tab1, om, tab2, num)):
            parent.insert(idx + k, el)
        pPr = prev.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr"); prev.insert(0, pPr)
        tabs = OxmlElement("w:tabs")
        for val, pos in (("center", int(usable.twips / 2)), ("right", int(usable.twips))):
            tb = OxmlElement("w:tab"); tb.set(qn("w:val"), val); tb.set(qn("w:pos"), str(pos)); tabs.append(tb)
        pPr.append(tabs)
        # 本文の行送りは固定18pt なので、式の段落だけ行送りを「1行」にして分数や和記号が切れないようにする
        sp = OxmlElement("w:spacing")
        for k, v in (("w:before", "120"), ("w:after", "120"), ("w:line", "240"), ("w:lineRule", "auto")):
            sp.set(qn(k), v)
        pPr.append(sp)
        ind = OxmlElement("w:ind"); ind.set(qn("w:firstLine"), "0"); pPr.append(ind)
        p.getparent().remove(p)


def page_breaks(doc):
    """改ページだけの段落を、次の段落の「段落前で改ページ」に置き換える（新しい頁の先頭に空行を残さない）。
    文書末の改ページだけの段落は、空白の最終頁を作るので削除する。"""
    body = doc.element.body
    BR = qn("w:br")
    for p in list(body.iterchildren(qn("w:p"))):
        runs = p.findall(qn("w:r"))
        if len(runs) != 1 or p.find(qn("w:pPr")) is not None:
            continue
        kids = list(runs[0])
        if len(kids) != 1 or kids[0].tag != BR or kids[0].get(qn("w:type")) != "page":
            continue
        nxt = p.getnext()
        while nxt is not None and nxt.tag in (qn("w:bookmarkStart"), qn("w:bookmarkEnd")):
            nxt = nxt.getnext()
        if nxt is None or nxt.tag == qn("w:sectPr"):
            body.remove(p)
        elif nxt.tag == qn("w:p"):
            Paragraph(nxt, doc._body).paragraph_format.page_break_before = True
            body.remove(p)


def page_field_paragraph(footer, fmt):
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for instr, kind in (("begin", None), (" PAGE ", "instr"), ("separate", None), ("1", "text"), ("end", None)):
        r = OxmlElement("w:r")
        if kind == "instr":
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr; r.append(it)
        elif kind == "text":
            tt = OxmlElement("w:t"); tt.text = instr; r.append(tt)
        else:
            fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), instr); r.append(fc)
        p._p.append(r)


def sections(doc):
    """〔SECTION-BREAK:…〕の段落を節区切りにし、頁番号（なし／i, ii／1, 2）を設定する"""
    body = doc.element.body
    final = body.find(qn("w:sectPr"))
    marks = []
    for p in list(body.iter(qn("w:p"))):
        text = "".join(t.text or "" for t in p.iter(qn("w:t")))
        m = re.fullmatch(r"\s*〔SECTION-BREAK:(\w+)〕\s*", text)
        if m:
            marks.append((m.group(1), p))
    for name, p in marks:
        for r in list(p):
            if r.tag != qn("w:pPr"):
                p.remove(r)
        pPr = p.find(qn("w:pPr"))
        if pPr is None:
            pPr = OxmlElement("w:pPr"); p.insert(0, pPr)
        sp = copy.deepcopy(final)
        for ref in sp.findall(qn("w:footerReference")) + sp.findall(qn("w:headerReference")):
            sp.remove(ref)
        pPr.append(sp)
    secs = doc.sections
    kinds = [m[0] for m in marks] + ["body"]
    for sec, kind in zip(secs, kinds):
        sp = sec._sectPr
        for old in sp.findall(qn("w:pgNumType")):
            sp.remove(old)
        pg = OxmlElement("w:pgNumType")
        if kind == "front":
            pg.set(qn("w:fmt"), "lowerRoman"); pg.set(qn("w:start"), "1")
        elif kind == "body":
            pg.set(qn("w:fmt"), "decimal"); pg.set(qn("w:start"), "1")
        sp.append(pg)
        sec.footer.is_linked_to_previous = False
        if kind != "cover":
            page_field_paragraph(sec.footer, kind)
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin, sec.bottom_margin = Mm(30), Mm(25)
        sec.left_margin, sec.right_margin = Mm(30), Mm(25)


def update_fields_on_open(doc):
    settings = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    # スキーマの順序（hdrShapeDefaults・footnotePr・compat・rsids などの前）に入れる
    later = {qn(f"w:{n}") for n in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids",
                                    "attachedSchema", "themeFontLang", "clrSchemeMapping", "doNotIncludeSubdocsInStats",
                                    "doNotAutoCompressPictures", "forceUpgrade", "captions", "readModeInkLockDown",
                                    "smartTagType", "schemaLibrary", "shapeDefaults", "doNotEmbedSmartTags",
                                    "decimalSymbol", "listSeparator")}
    later.add("{http://schemas.openxmlformats.org/officeDocument/2006/math}mathPr")
    nxt = [c for c in settings if c.tag in later]
    if nxt:
        nxt[0].addprevious(uf)
    else:
        settings.append(uf)


def convert(md, out, ref):
    pypandoc.convert_text(md, "docx", format="markdown+tex_math_dollars+pipe_tables+footnotes+raw_attribute+implicit_figures+fenced_divs",
                          outputfile=out, extra_args=[f"--reference-doc={ref}", f"--resource-path={HERE}"])


def postprocess(path, with_sections=True):
    doc = Document(path)
    three_line_tables(doc)
    number_equations(doc)
    page_breaks(doc)
    for p in doc.paragraphs:  # 注）・出所）の段落を小さく
        if re.match(r"^\s*(注|出所)）", p.text):
            p.style = doc.styles["Table Note"]
    if with_sections:
        sections(doc)
    update_fields_on_open(doc)
    doc.save(path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--tag", default="draft")
    a = ap.parse_args()
    main_files, sep_files = collect(a.tag)
    md = assemble(main_files)
    open(os.path.join(HERE, f"thesis_{a.tag}.md"), "w", encoding="utf-8").write(md)
    with tempfile.TemporaryDirectory() as td:
        ref = os.path.join(td, "reference.docx")
        make_reference(ref)
        out = os.path.join(HERE, f"thesis_{a.tag}.docx")
        convert(md, out, ref)
        postprocess(out)
        print("本体:", out)
        if sep_files:
            smd = fix_math("\n\n".join(open(f, encoding="utf-8").read() for f in sep_files))
            sout = os.path.join(HERE, f"abstract_{a.tag}.docx")
            convert(smd, sout, ref)
            postprocess(sout, with_sections=False)
            print("別冊要旨:", sout)


if __name__ == "__main__":
    main()
